# -*- coding: utf-8 -*-
"""research_discussion_transcript.md → docx 변환.
헤딩(#/##/###), 코드펜스(```), 표(| |), 불릿(-), 구분선(---), 굵게(**bold**) 처리.
사용: python scripts/md_to_docx_transcript.py
"""
import os, re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

_HERE = os.path.dirname(os.path.abspath(__file__))
_REPO = os.path.dirname(_HERE)
SRC = os.path.join(_REPO, "docs", "research_discussion_transcript.md")
DST = os.path.join(_REPO, "docs", "research_discussion_transcript.docx")

BOLD = re.compile(r"\*\*(.+?)\*\*")


def add_runs(p, text):
    """**bold** 인라인 처리."""
    pos = 0
    for m in BOLD.finditer(text):
        if m.start() > pos:
            p.add_run(text[pos:m.start()])
        r = p.add_run(m.group(1))
        r.bold = True
        pos = m.end()
    if pos < len(text):
        p.add_run(text[pos:])


def main():
    lines = open(SRC, encoding="utf-8").read().split("\n")
    doc = Document()
    doc.styles["Normal"].font.name = "Malgun Gothic"
    doc.styles["Normal"].font.size = Pt(10)

    i = 0
    in_code = False
    code_buf = []
    while i < len(lines):
        line = lines[i]

        # 코드펜스
        if line.strip().startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                run = p.add_run("\n".join(code_buf))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                p.paragraph_format.left_indent = Pt(12)
                code_buf = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        # 표: 연속된 | 라인 묶기
        if line.strip().startswith("|"):
            tbl_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                tbl_lines.append(lines[i].strip())
                i += 1
            rows = []
            for tl in tbl_lines:
                if re.match(r"^\|[\s:|-]+\|?$", tl):  # 구분선 행
                    continue
                cells = [c.strip() for c in tl.strip("|").split("|")]
                rows.append(cells)
            if rows:
                ncol = max(len(r) for r in rows)
                t = doc.add_table(rows=0, cols=ncol)
                t.style = "Light Grid Accent 1"
                for r in rows:
                    cells = t.add_row().cells
                    for j in range(ncol):
                        cells[j].text = r[j] if j < len(r) else ""
            continue

        stripped = line.strip()
        if stripped == "---":
            doc.add_paragraph().add_run("").add_break()
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, stripped[2:])
        elif stripped == "":
            pass
        else:
            p = doc.add_paragraph()
            add_runs(p, line)
        i += 1

    doc.save(DST)
    print("wrote", DST)


if __name__ == "__main__":
    main()
