# -*- coding: utf-8 -*-
"""실험 상세 정리 리포트(.docx) 생성.

구성: 0.전체 개요 → 1.실험 설계 → 2.실험 과정 → 3.결과 → 4.한계·향후 → 5.부록
수치는 본실험(Qwen3.8-27B-FP8, n=100) 확정값. 재생성: python scripts/build_experiment_report.py
출력: docs/experiment_report.docx
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
OUT = os.path.join(ROOT, "docs", "experiment_report.docx")

doc = Document()
# 기본 폰트
st = doc.styles["Normal"]
st.font.name = "맑은 고딕"
st.font.size = Pt(10.5)


def h(text, level=1):
    doc.add_heading(text, level=level)


def p(text="", bold=False, italic=False):
    par = doc.add_paragraph()
    run = par.add_run(text)
    run.bold = bold
    run.italic = italic
    return par


def bullet(text):
    doc.add_paragraph(text, style="List Bullet")


def formula(text):
    """수식 한 줄(들여쓰기·이탤릭)."""
    par = doc.add_paragraph()
    par.paragraph_format.left_indent = Pt(24)
    run = par.add_run(text)
    run.italic = True
    run.font.name = "Cambria Math"
    return par


def defn(text):
    """기호 정의 한 줄(작은 들여쓰기)."""
    par = doc.add_paragraph()
    par.paragraph_format.left_indent = Pt(24)
    run = par.add_run(text)
    run.font.size = Pt(9.5)
    return par


def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, hd in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = str(hd)
        for r in c.paragraphs[0].runs:
            r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    return t


# ── 표지 ────────────────────────────────────────────────
title = doc.add_heading("로컬 오픈소스 LLM과 방언 지식 주입을 활용한\n"
                        "저자원 한국어 방언 악취 민원 대화 합성 — 실험 상세 정리", 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p("본 문서는 논문 Ⅳ장(실험·결과) 집필을 위한 상세 정리본이다. 개요 → 실험 설계 → "
  "실험 과정 → 결과 → 한계 순으로 기술한다.").italic = True

# ── 0. 전체 개요 ─────────────────────────────────────────
h("0. 전체 개요 (Outline)", 1)

h("0.1 연구 목표", 2)
p("상용 API·파인튜닝 없이 로컬 오픈소스 LLM만으로 5개 권역(경상·전라·충청·강원·제주)의 "
  "방언 악취 민원 상담 대화를 합성하고, 방언 지식을 어떤 방식으로 공급하는 것이 효과적인지를 "
  "정직하게 비교·진단한다.")

h("0.2 기여", 2)
p("주 기여 — 로컬 LLM 기반 5권역 방언 악취 민원 대화 합성 파이프라인(S1/S2)과 그 산출 데이터셋.", bold=True)
p("부차 기여 — (1) 코퍼스 유래 마커 기반 reference-free 방언성 자동 진단 지표와 독립 모델 "
  "교란을 통한 신뢰성 검증, (2) few-shot과 지식주입(RAG)의 효과에 대한 정직한 비교 관찰.")
p("주의: 본 연구는 'RAG가 우월하다'를 주장하지 않는다. 모든 지표는 대리 지표(proxy)이며 "
  "원어민 평가가 부재하므로, 어떤 방식의 절대적 우월성도 단정하지 않는다.", italic=True)

h("0.3 실험 한눈에", 2)
table(["평가 축", "무엇을 재나", "성격", "핵심 결과"],
      [["어미 방언율", "종결어미 방언 실현", "주 결과(독립·견고)", "few-shot 우위"],
       ["자연스러움(LLM judge)", "holistic 방언 품질", "품질축(독립·견고)", "few-shot 우위, S2 대등"],
       ["지표 신뢰성(A단계)", "지표가 오류에 반응하나", "타당성 검증", "민감도.87/특이도.93"],
       ["주입 활용률", "준 지식을 실제 쓰나", "메커니즘", "어휘21%/규칙11%"]])

h("0.4 핵심 결론 (3줄)", 2)
bullet("few-shot(완결 예시)이 방언 실현·자연스러움에서 지식주입(RAG)보다 효과적이다.")
bullet("RAG의 기여는 어휘 중심이며 S2(2단계 변환)에 한정되어 few-shot과 대등한 수준에서 그친다.")
bullet("어미 카운팅과 holistic judge 두 독립 평가가 동일 방향으로 수렴한다.")

# ── 1. 실험 설계 ─────────────────────────────────────────
h("1. 실험 설계", 1)

h("1.1 모델 및 환경", 2)
p("로컬 vLLM 서버에 배치한 Qwen3.8-27B-FP8 단일 모델로 모든 대화를 생성하였다. 상용 API와 "
  "파인튜닝은 사용하지 않았다. FP8은 로컬 GPU 서빙이 가능한 양자화 형식이며, 이전 예비 실험의 "
  "모델 혼재(3.6/3.8)를 제거하기 위해 본실험은 단일 모델로 통일하였다.")

h("1.2 시나리오", 2)
p("각 권역에 대해 고정 시드로 악취 민원 시나리오를 표집하였다(총 n=100). 입력은 구조화된 "
  "키워드 슬롯(냄새 종류, 민원인 위치, 냄새 강도 변화, 지속 시간, 원인 추정 지역)이며, 출력은 "
  "상담원(표준어)과 민원인(해당 권역 방언) 사이의 8~10턴 상담 대화이다. 모든 조건·파이프라인에 "
  "동일 시나리오를 적용해 공정 비교하였다.")

h("1.3 지식 공급 조건", 2)
table(["조건", "공급 지식", "분리 대상"],
      [["C0", "없음(금지 규칙만)", "기준선(zero-shot)"],
       ["C1", "고정 few-shot 예시", "예시 제공 효과"],
       ["C_kb", "RAG(어휘+문법 규칙+예문)", "검색·주입 지식 효과"],
       ["C_kb+C1", "RAG + few-shot 예시", "결합 효과"]])
p("C1은 완결된 방언 대화 예시를 제시하는 필수 대조군으로, 품질 향상이 '지식 검색' 때문인지 "
  "단지 '예시가 제공되어서'인지를 분리한다.")

h("1.4 생성 파이프라인", 2)
p("S1(직접 생성): 키워드로부터 방언 대화를 한 단계로 생성한다. 생성 이전에 질의로 삼을 문장이 "
  "없어, 어미 지식은 권역 조건부 고정 주입이 지배한다(주입, injection).")
p("S2(2단계 변환): 표준어 대화를 먼저 생성한 뒤, 생성된 표준어 발화를 질의로 방언 지식을 "
  "검색·주입하여 방언으로 변환한다. 질의가 완결 문장이라 어휘 검색이 활발하다(검색, retrieval).")

h("1.5 RAG 구성 — 어휘 검색 vs 어미 고정주입", 2)
p("방언 지식은 두 층위로 나뉜다. (1) 어휘: 표준어 표제어를 정확 매칭하여 방언형을 조회하며, "
  "질의마다 결과가 달라지는 진짜 질의조건부 검색이다. (2) 어미/문법: 권역별 종결어미 등은 발화 "
  "주제와 독립적이라 검색 대상이 되기 어려워, 해당 범주를 질의와 무관하게 고정 주입한다. 즉 "
  "본 시스템은 '어휘 검색 + 어미 주입'의 하이브리드이며, 정교한 검색이라 과대포장하지 않는다.")

h("1.6 평가 지표 설계", 2)
bullet("어미 방언율(주): 민원인 발화의 종결어미가 권역 방언 어미인지. 마커는 AI-Hub 코퍼스에서 "
       "추출(독립). S1은 표면 접미 매칭, S2는 표준-방언 정렬쌍 비교(reference-based).")
bullet("자연스러움(LLM judge): 표면 카운팅이 못 잡는 holistic 방언 품질을 독립 모델(Opus) "
       "블라인드 쌍대비교로 평가.")
bullet("지표 신뢰성(A단계): 채점기와 독립인 참조 모델로 정상문을 표준어화·타권역화해 지표 반응 측정.")
bullet("주입 활용률(메커니즘): 주입한 어휘·규칙 중 생성문에 실제 등장한 비율.")

h("1.7 지표 계산식 (수식)", 2)

p("가. 어미 방언율·누수율", bold=True)
defn("y : 민원인 발화, W(y) : y의 어절 집합, |W| : 어절 수")
defn("M⁺_r : 권역 r의 방언 마커 집합, M⁻_r : 타권역 마커 집합, 1[·] : 지시함수")
p("S1(직접 생성) — 표면 접미 매칭. 어절 w가 마커 t로 끝나는지(길이·받침 가드 적용):")
formula("plus  = Σ_{w∈W} 1[ ∃ t∈M⁺_r : w ⊇suffix t ]")
formula("minus = Σ_{w∈W} 1[ ∃ t∈M⁻_r : w ⊇suffix t ]")
p("S2(2단계 변환) — 표준·방언 어절을 정렬해 변환쌍만 채점. 공통 어간 이후 꼬리 (s_t, d_t):")
formula("P = alignedpairs(std, dia),   (s_t, d_t) = tail(s, d)")
formula("plus = Σ_{(s,d)∈P} 1[ d_t ∈ M⁺_r ],   minus = Σ 1[ d_t ∈ M⁻_r ]")
p("공통(N = |W| 또는 |P|):")
formula("방언율  D⁺(y,r) = plus / N")
formula("누수율  D⁻(y,r) = minus / N")
formula("결합지표 D(y,r) = ( plus − minus ) / N")

p("나. 지표 신뢰성 (A단계 교란)", bold=True)
defn("ŷ : 정상 생성문, ỹ_std : 표준어화(방언 제거), ỹ_cross : 타권역화")
formula("민감도    Sensitivity = ( D⁺(ŷ,r) − D⁺(ỹ_std, r) ) / D⁺(ŷ,r)")
formula("특이도-A  SrcDrop   = ( D⁺(ŷ,r) − D⁺(ỹ_cross, r) ) / D⁺(ŷ,r)")
formula("특이도-B  LeakRise  = D⁻(ỹ_cross, r) − D⁻(ŷ, r)")
defn("1에 가까울수록(민감도·특이도-A) 지표가 방언 제거·타권역 오염을 잘 탐지.")

p("다. 주입 활용률", bold=True)
defn("I_v, I_r : 재현한 주입 어휘·규칙 형태 집합, y_kb : C_kb 생성문, y_0 : C0 생성문")
formula("어휘 활용률  U_v = | { f∈I_v : f ∈ y_kb } | / |I_v|")
formula("규칙 활용률  U_r = | { f∈I_r : f ∈ y_kb } | / |I_r|")
formula("귀속 상승분  A_v = | { f∈I_v : f∈y_kb ∧ f∉y_0 } | / |I_v|")

p("라. LLM judge (자연스러움)", bold=True)
defn("개별 판정엔 명시적 채점식이 없다 — 독립 모델의 블라인드 쌍대 선택이 판정 함수이다.")
defn("J_i(A,B) ∈ {A, B} : i번째 시나리오에서 judge가 더 자연스럽다고 고른 조건(순서 무작위)")
formula("승률  Win(A≻B) = (1/N) Σ_{i=1}^{N} 1[ J_i(A,B) = A ]")
formula("신뢰성 정답률  Acc = (1/N) Σ_i 1[ J_i(원본, 교란) = 원본 ]")
defn("Win(A≻B) > 0.5 이고 유의하면 A가 더 자연스러움. Acc가 높으면 judge 신뢰(관문).")

# ── 2. 실험 과정 ─────────────────────────────────────────
h("2. 실험 과정", 1)

h("2.1 대화 생성", 2)
p("Qwen3.8-27B-FP8로 n=100 시나리오에 대해 S1·S2 각각 4조건을 생성하였다. 어휘 검색 coverage "
  "평균은 S1 0.242, S2 0.302였다.")

h("2.2 어미 방언율 채점 및 통계", 2)
p("민원인 발화만 추출해 어절 단위로 정규화한 뒤, S1은 어절의 종결부가 방언 마커로 끝나는지, "
  "S2는 표준-방언 정렬쌍의 꼬리 변화가 방언 마커인지로 채점하였다. 타깃 권역 마커는 방언율(+), "
  "타권역 마커는 누수율(−)로 집계하였다. 통계는 시나리오 짝지은 부트스트랩(95% CI, 양측 p)로 "
  "수행하였다.")

h("2.3 방법론 점검 — 어휘축 탐색과 제외 (정직성)", 2)
p("어미 방언율은 권역 고정 어미를 재므로 RAG의 검색 파트(어휘)를 포착하지 못한다. 이에 어휘 "
  "실현을 별도 측정하였으나, 다음 과정을 거쳐 본 지표를 본문 결과에서 제외하였다.")
bullet("초기 어휘 지표의 참조가 RAG 검색 사전과 동일해, C_kb에 구조적으로 유리한 '홈 어드밴티지' "
       "편향이 있음을 확인하였다.")
bullet("AI-Hub 코퍼스에서 독립적으로 방언 어휘를 추출(표준-방언 정렬쌍의 어간 변화)해 재검증한 "
       "결과, 참조 구성(용언 포함/명사 한정)에 따라 방향이 뒤집혔고 절대값도 매우 작았다.")
bullet("따라서 어휘축 우위는 견고하지 않다고 판단하고, 판단 근거로 삼지 않는다. 다만 'RAG의 검색 "
       "기여는 본질적으로 어휘'라는 개념적 통찰은 유지한다.")
p("이 과정은 '원하는 결과가 나올 때까지 지표를 바꾸는 것(metric shopping)'을 피하기 위한 것으로, "
  "지표 청소는 결과와 무관한 객관 규칙(형태소 분석기 표준어 제외 등)으로만 수행하였다.", italic=True)

h("2.4 LLM judge 도입 — 신뢰성 관문 후 조건 평가", 2)
p("원리: LLM judge는 개별 텍스트에 절대 점수를 매기지 않고, 두 발화를 블라인드로 제시해 '어느 "
  "쪽이 더 자연스러운 방언인가'를 이진 선택하게 한다. 명시적 채점 수식 없이 독립 모델의 암묵적 "
  "판단을 인간 평가의 대리(proxy)로 사용하며, 절대 점수(예: 1~5)보다 보정·안정성이 높은 쌍대비교를 "
  "택하였다. 제시 순서를 무작위화해 위치 편향을 통제하고, judge(Opus)와 생성 모델(Qwen)이 달라 "
  "self-preference를 배제한다. 판정은 승률로 집계한다(식은 1.7 라 참조).")
p("먼저 judge의 신뢰성을 검증하였다. A단계 교란 산출물(원본 방언 / 표준어화 / 타권역화)을 "
  "재활용해, 원본이 교란본보다 더 자연스러운 방언으로 선택되는 정답률을 측정하였다(정답을 아는 "
  "대비이므로 관문 역할). 이어 검증된 judge를 실제 조건(C_kb vs C1 등) 비교에 적용하였다.")

h("2.5 지표 신뢰성 교란 (A단계)", 2)
p("채점기(마커)와 독립인 Opus로 정상 방언문을 (a) 표준어화(방언 제거)·(b) 타권역화(다른 권역 "
  "방언으로 변환)하여, 지표가 방언 제거에 민감하게 반응하고(민감도) 타권역 오염을 분별하는지"
  "(특이도)를 측정하였다.")

h("2.6 주입 활용률 분석", 2)
p("검색은 결정적이므로 저장된 질의로 주입 목록을 재현하고, C_kb 생성문에 주입된 어휘·규칙이 "
  "실제 등장한 비율을 집계하였다(LLM 호출 없이 계산).")

# ── 3. 결과 ─────────────────────────────────────────────
h("3. 결과", 1)

h("3.1 어미 방언율 (주 결과)", 2)
table(["조건", "S1", "S2"],
      [["C0", "0.014", "0.048"],
       ["C_kb", "0.050", "0.060"],
       ["C1", "0.084", "0.199"],
       ["C_kb+C1", "0.085", "0.112"]])
p("주요 짝지은 대비(부트스트랩): S1에서 C1−C_kb = +0.034(p<.001), C1−C_kb+C1 = −0.001(ns). "
  "S2에서 C_kb−C0 = +0.012(ns; RAG 단독은 기준선도 유의하게 못 넘음), C1−C_kb+C1 = "
  "+0.087(p<.001; few-shot에 RAG를 더하면 오히려 하락).")
p("→ few-shot(C1)이 어미 실현을 지배하며, 지식주입(C_kb)의 고유 기여는 유의하지 않거나 음(−)이다.", bold=True)

h("3.2 자연스러움 (LLM judge)", 2)
p("Judge 신뢰성: 원본 vs 표준어화 98.9%, 원본 vs 타권역화 98.9%로 원본을 선택 → 방언·권역 "
  "분별 확인(명확한 대비에 대한 필요조건).")
table(["비교", "S1", "S2"],
      [["C1 > C0 (앵커)", "97%", "89%"],
       ["C_kb > C1 (핵심)", "20%", "43%"]])
p("앵커(C1>C0)가 높아 우리 데이터에서도 judge가 타당함을 확인하였다. 핵심 대비 C_kb>C1은 "
  "S1에서 20%(few-shot 우세), S2에서 43%(50%와 유의차 없음, 사실상 대등)였다.")
p("→ RAG는 자연스러움에서 few-shot을 넘지 못하며, S2에서만 대등한 수준에 이른다.", bold=True)

h("3.3 지표 신뢰성 (A단계)", 2)
p("n=88(S1·C_kb+C1). 민감도 0.872(표준어화 시 방언율 하락), 특이도 0.931(타권역화 시 원권역 "
  "방언율 하락), 누수율 상승 +0.113, 변환 확인 0.122. 지표가 통제된 오류에 정확히 반응하여 "
  "진단 도구로서 신뢰성을 확인하였다.")

h("3.4 주입 활용률 (메커니즘)", 2)
p("C_kb에서 주입된 지식 중 생성문 등장 비율은 어휘 21%, 규칙 11%였으며, 어휘가 전혀 검색되지 "
  "않은 시나리오가 28/100이었다. 모델이 단편적 주입 지식을 충분히 적용하지 않는다는 점이, "
  "완결 예시를 제시하는 few-shot이 어미 실현에서 우세한 관찰된 원인이다.")

h("3.5 종합 — 이중검증 수렴", 2)
p("어미 카운팅(3.1)과 holistic judge(3.2) 두 독립 평가가 동일 방향으로 수렴한다: few-shot이 "
  "방언 실현·자연스러움에서 더 효과적이며, 지식주입(RAG)의 기여는 어휘 중심·S2 한정으로 few-shot과 "
  "대등한 수준에서 그친다. 어느 방식도 일률적으로 우월하지 않고, RAG가 few-shot을 능가한다는 근거는 "
  "확인되지 않았다.")

# ── 4. 한계 및 향후 ──────────────────────────────────────
h("4. 한계 및 향후 과제", 1)
bullet("원어민 평가 부재 — 모든 지표는 대리 지표. 절대적 품질 우월은 단정 불가.")
bullet("어휘축은 참조 민감·결론 불가로 본문에서 제외(2.3).")
bullet("강원·제주는 방언 실현 자체가 약해 검증 신뢰도가 낮음.")
bullet("RAG 문법 규칙 검수는 일부 권역만 수행.")
bullet("threshold(모델 크기별 RAG 효과)는 서버 제약으로 미실행 — 향후 과제.")
bullet("향후: 원어민/추가 LLM judge, 모델 크기 사다리(threshold), 코퍼스 어휘 보강 후 judge 재평가.")

# ── 5. 부록 ─────────────────────────────────────────────
h("5. 부록", 1)
h("5.1 어미 방언율·누수율 전체 (조건×파이프라인)", 2)
p("(집계 파일: results/eval_conditions.csv · 통계: src/eval/stats_test.py)")
h("5.2 주요 산출 파일", 2)
bullet("생성: results/s1_20260818_115333.jsonl, results/s2_20260818_135335.jsonl")
bullet("어미 통계: src/eval/stats_test.py")
bullet("어휘 독립참조(내부 보존): src/eval/extract_vocab_pairs_aihub.py, docs/dialect_vocab_aihub.csv")
bullet("judge: src/eval/judge_reliability.py, judge_conditions.py")
bullet("A단계: src/eval/llm_perturb_sensitivity.py")
bullet("활용률: src/eval/injection_utilization.py")

doc.save(OUT)
print("저장:", os.path.relpath(OUT, ROOT))
