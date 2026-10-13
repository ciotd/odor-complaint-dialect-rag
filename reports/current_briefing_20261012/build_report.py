"""Current-state Korean briefing. Read frozen research artifacts; never call models."""
import ast
from collections import Counter
import hashlib
import json
from pathlib import Path
import subprocess
import sys

sys.path.insert(0, '/tmp/kg-briefing-deps')
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
EXP = 'reports/kg_v2_factorial_20261012/'
ADV = 'reports/kg_v2_advisory_20261012/'
used = set()

def read(path):
    used.add(path)
    return json.loads((ROOT/path).read_text())

stats = read(ADV+'retrieval_statistics.json')
build = read('kg_v2/reports/build_summary.json')
registry = read('kg_v2/source_registry.json')
graph = read('kg_v2/graph.json')
protocol = read(EXP+'protocol.json')
summary = read(EXP+'summary.json')
metrics = read(EXP+'review_metrics.json')
jobs = read(EXP+'jobs.json')
outputs = read(EXP+'outputs.json')
analyses = read(EXP+'analysis.json')['items']
retrieval = read(EXP+'retrieval.json')
peer = read('reports/model_comparison_20261012/comparison.json')
peer_protocol = read('reports/astra_high_no_kg_20261012/protocol.json')
advisor = read(ADV+'advisor.json')
critic = read(ADV+'counterexamples.json')
literature = read(ADV+'literature.json')
audit = read(ADV+'coordinator_validation.json')
reviewed = read(ADV+'reviewed_advice.json')
seg_rows = read(ADV+'segment_lengths.json')
out_by = {(r['id'],r['condition']):r['output'] for r in outputs}
input_by = {r['id']:r['input'] for r in analyses}
by_source = {s['source_id']:s for s in registry['sources']}

doc = Document()
sec = doc.sections[0]
sec.page_width,sec.page_height = Cm(21),Cm(29.7)
sec.top_margin,sec.bottom_margin = Cm(1.8),Cm(1.7)
sec.left_margin=sec.right_margin=Cm(1.9)
sec.header_distance=sec.footer_distance=Cm(.7)
for name in ['Normal','Body Text','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    style=doc.styles[name]
    style.font.name='Noto Sans CJK KR'
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Noto Sans CJK KR')
doc.styles['Normal'].font.size=Pt(10)
doc.styles['Normal'].paragraph_format.line_spacing=1.15
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
for name,size in [('Title',28),('Heading 1',18),('Heading 2',12),('Heading 3',10.5)]:
    s=doc.styles[name];s.font.size=Pt(size);s.font.color.rgb=RGBColor.from_string('163B55')
    s.paragraph_format.keep_with_next=True
doc.styles['Heading 1'].paragraph_format.page_break_before=True
doc.styles['Caption'].font.size=Pt(8)
doc.styles['Caption'].font.color.rgb=RGBColor.from_string('596775')
sec.header.paragraphs[0].text='경상도 방언 KG 연구 · 현재 구현 및 최신 실험 보고'
sec.header.paragraphs[0].style='Caption'
f=sec.footer.paragraphs[0];f.alignment=2
f.add_run('2026-10-12 기준  |  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
doc.core_properties.title='경상도 방언 KG 연구 — 현재 구현·실험·개선 과제 상세 보고'
doc.core_properties.subject='2026-10-13 보고용; KG v2, 50개 입력 A/B/C/D, 실제 전달량, 4역할 검토'
doc.core_properties.author='연구 프로젝트 기록 정리'

def p(t='',style=None):return doc.add_paragraph(str(t),style)
def body(t):
    for line in t.strip().split('\n\n'):p(line.strip())
def h(t,level=1):
    par=doc.add_heading(t,level)
    name='s'+hashlib.md5(t.encode()).hexdigest()[:12]
    a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(len(doc.paragraphs)));a.set(qn('w:name'),name)
    b=OxmlElement('w:bookmarkEnd');b.set(qn('w:id'),str(len(doc.paragraphs)))
    par._p.insert(0,a);par._p.append(b)
    return name
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1';t.autofit=False
    if widths:
        for c,w in zip(t.columns,widths):c.width=Cm(w)
    for c,v in zip(t.rows[0].cells,headers):c.text=str(v)
    rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
    for values in rows:
        cells=t.add_row().cells
        for c,v in zip(cells,values):c.text=str(v)
    for row in t.rows:
        ns=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(ns)
        for i,c in enumerate(row.cells):
            if widths:c.width=Cm(widths[i])
            for pp in c.paragraphs:
                pp.paragraph_format.space_after=Pt(3)
                pp.paragraph_format.line_spacing=1.05
                for r in pp.runs:r.font.size=Pt(8.5)
    p()
    return t
def code(value):
    text=value if isinstance(value,str) else json.dumps(value,ensure_ascii=False,indent=2)
    for line in text.splitlines():
        pp=p(line or ' ');pp.paragraph_format.space_after=Pt(0);pp.paragraph_format.line_spacing=1.05
        for r in pp.runs:r.font.size=Pt(8)
    p()
def caption(t):p(t,'Caption')
def figure(name,dot,caption_text):
    (OUT/(name+'.dot')).write_text(dot)
    subprocess.run(['dot','-Tpng','-Gdpi=180',str(OUT/(name+'.dot')),'-o',str(OUT/(name+'.png'))],check=True)
    pp=p();pp.alignment=1;pp.add_run().add_picture(str(OUT/(name+'.png')),width=Cm(17))
    caption(caption_text)
def citation(t):caption('근거: '+t)
def output_case(sid,comment):
    h(sid+' · '+input_by[sid],2)
    table(['조건','저장된 실제 번역'],[(c,out_by[sid,c]) for c in 'ABCD'],[1.4,15.8])
    p(comment)

chapters=[
'1. 보고 요지와 현재 단계','2. 용어와 시스템 경계','3. 전체 처리 과정과 모델 역할',
'4. KG v2의 문헌·규모·구조','5. 입력 분석과 기본 조회 방식','6. 의미 검색과 구간별 변환 계획',
'7. 최신 A/B/C/D 실험 설계와 실행','8. 실제 전달량: 정의·계산·결과','9. 번역 결과와 해석',
'10. 네 역할의 검토와 원문 대조','11. 확인된 원인과 남은 가설','12. 수정 순서와 검증 설계',
'13. 보고용 설명과 예상 질의응답','부록 A. 실제 프롬프트와 데이터 계약',
'부록 B. 최신 50개 입력의 A/B/C/D 출력','부록 C. 새 대조 입력 20개와 조언 9건',
'부록 D. 근거 파일·재현·검증 범위']

p('경상도 방언 KG 연구','Title')
p('현재 구현 · 최신 실험 · 개선 과제','Subtitle')
p('상사 보고용 상세 자료 | 보고 예정일 2026년 10월 13일')
p('자료 기준일: 2026년 10월 12일, 네 역할 검토·통합 완료 시점')
body('''핵심 질문: 문헌의 방언 지식을 어떤 단위로 표현하고, 어떤 조건으로 조회하여, 생성기에 어떻게 전달해야 원문의 의미와 높임을 유지한 방언 생성으로 이어지는가?

현재 성과는 문헌 근거를 추적할 수 있는 KG v2, 분석·조회·생성의 연결, 50개 입력에 대한 네 조건 200개 출력, 실제 전달량 집계, 역할을 나눈 원인 검토다. 전체 번역 품질의 우월성을 확정한 단계는 아니다.

이번 보고서는 현재 상태를 중심으로 작성했다. 오래된 버전의 세부 이력은 줄이고, 기본 실행 경로·분리된 실험 경로·아직 구현하지 않은 보완안을 구별한다. 본문은 판단과 구조, 부록은 실제 프롬프트·전체 출력·후속 진단 자료를 제공한다.''')
table(['현재 기준','확인된 내용'],[
('지식 자료','7종·474쪽 / 3,761개 지식 진술 / 25,783개 노드 / 35,358개 관계'),
('최신 생성 실험','기존 38개 + 추가 12개 = 50개 입력 × A/B/C/D = 200개 출력'),
('최신 원인 검토','별도 작업자 3명 + 조정자 1명 / 제안 9건 / 원문 인용 검증 22건'),
('현재 반영 경계','조언자 검토 결과는 별도 자료. 이번 검토에서 기본 KG·조회기·생성기는 변경하지 않음')],[3.4,13.8])
p('읽는 순서: 1장과 13장으로 보고 흐름을 잡고, 2~7장으로 현재 구조를 확인한 뒤, 8~12장으로 수치·문제·다음 조치를 검토한다.','Caption')

h('목차')
for title in chapters:
    pp=p();link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),'s'+hashlib.md5(title.encode()).hexdigest()[:12])
    run=OxmlElement('w:r');text=OxmlElement('w:t');text.text=title;run.append(text);link.append(run);pp._p.append(link)
p('목차 항목을 클릭하면 해당 절로 이동한다. 쪽 번호는 Word의 페이지 필드로 표시한다.','Caption')

h(chapters[0])
body('''연구 대상은 일반 표준어 발화를 경상도 방언으로 바꾸는 과제다. 의미·시제·부정·높임·질문과 부탁 같은 발화 의도를 유지하면서 자연스러운 방언 표현을 생성하는 것이 목표다. 방언 단어가 많아지거나 표준어와 달라지는 것 자체를 정확도로 보지 않는다.

현재 KG v2에는 문헌의 어휘 대응뿐 아니라 활용형, 종결어미, 구문, 용법, 제약, 출처와 원문 근거가 들어 있다. 입력을 Qwen과 Kiwi로 분석하고, 조건에 맞는 지식을 Python 조회기로 찾아, Qwen에 참고자료로 전달하는 경로까지 연결되어 있다.

최신 실험에서는 같은 KG와 Qwen 설정을 유지하면서 조회 방식과 전달 형식을 각각 바꿨다. B/D에는 의미 검색을 추가했고, C/D에는 구간별 변환 계획을 사용했다. 50개 입력 중 37개에서 후보 집합이 달라졌지만, 그 변화가 번역 품질의 일관된 향상으로 이어졌다고 결론 내릴 수는 없다.

최근 검토는 문제를 더 구체적으로 나누었다. 원문에 있는데 추출하지 못한 예문, KG에는 있는데 조회되지 않는 보존 정보, 검색 대상 자체가 만들어지지 않는 파생 서술어, 목록의 형태를 실제 변환 규칙처럼 쓰는 문제, 올바른 근거를 전달해도 잘못 생성하는 문제를 각각 조사해야 한다.

모델 조언을 이용한 보완도 시작했다. 변환 조언자가 규칙 가설을 제안하고, 반례 검토자가 과잉 일반화를 지적하며, 문헌 대조자가 실제 근거를 확인하고, 조정자가 통합했다. 이 자료는 문헌 사실과 구분하여 저장했으며 아직 기본 실행 코드에 적용하지 않았다.''')
table(['보고 가능한 결론','아직 보고하면 안 되는 결론'],[
('200개 출력을 생성하고 요청·응답·조회 기록을 보존했다.','KG 방식이 강한 사전학습 모델보다 우수하다.'),
('입력당 후보 전달 구간은 평균 3.0~3.2개, 구간당 후보는 약 2.75개다.','후보 수 또는 길이를 늘리거나 줄이면 반드시 개선된다.'),
('일부 분석·조회 누락과 추출 누락을 구체적 사례로 확인했다.','전체 문헌의 지식을 완전하게 추출했다.'),
('원문 대조로 기존 오류 판정 두 가지를 정정했다.','방언 화자 독립 평가와 통계적 우열 검증이 끝났다.')],[8.6,8.6])
citation('kg_v2/reports/build_summary.json; factorial protocol/summary; advisory coordinator_validation.json')

h(chapters[1])
table(['용어','현재 프로젝트에서의 의미'],[
('지식 진술 / KnowledgeStatement','문헌에서 읽은 하나의 대응·활용·용법·제약을 출처와 조건을 붙여 기록한 단위. 진술 1개가 반드시 독립적인 완성 변환 규칙 1개는 아니다.'),
('노드 / 관계','노드는 진술·표현·조건·근거 등의 객체다. 관계는 어떤 진술이 어떤 표현·조건·근거와 연결되는지를 나타낸다.'),
('카드 / 참조 후보','KG 노드 유형이 아니다. 조회기가 하나의 진술과 관련 정보, 원문 구간, 검사 결과를 묶어 생성기에 보내는 JSON 객체를 뜻한다.'),
('입력 / 발화','DEV001 같은 ID 하나에 해당하는 표준어 텍스트. LONG 입력에는 문장 여러 개가 들어가지만 실험에서는 입력 하나다.'),
('분석 절 / clause','분석기가 주절 또는 연결절로 기록한 단위. 통계의 후보 전달 구간과 같지 않다.'),
('원문 구성 / source_unit','인용·관형절·보조용언 같은 구성과 부모/자식 관계를 표현하는 분석 단위. 분석에 기록됐다고 모든 단위가 현행 v2 조회에서 활용되는 것은 아니다.'),
('구간 / span','원문에서 시작 위치와 끝 위치로 지정한 범위. [start,end)는 끝 위치를 포함하지 않는다. 단어 구간이 절 구간 안에 들어갈 수 있다.'),
('후보','해당 구간에 참고할 수 있다고 선택된 진술. 자동 승인된 정답이나 반드시 적용해야 할 치환이 아니다.'),
('청자 높임 / 주체 높임','듣는 사람에 대한 말투와 문장 속 행위 주체를 높이는 표현을 각각 뜻한다. -요와 -시-/계시다를 같은 변수로 처리하면 안 된다.'),
('벡터 검색 / 의미 검색','질의와 지식 설명을 숫자 벡터로 바꾸어 유사도로 후보를 찾는 방식. 유사성은 적용 가능성의 충분조건이 아니다.'),
('변환 계획','이번 C/D에서는 같은 구간의 후보와 보존할 분석값을 묶은 생성 입력. 모든 후보의 완성 활용형을 확정하는 규칙 실행기는 아니다.'),
('생산 KG·생성기','대화에서 실제 실행에 사용 중인 기본 데이터·코드를 가리킨 표현. 외부 상용 서비스에 배포했다는 뜻으로 쓰지 않는다.')],[4.1,13.1])
h('2.1 “변경하지 않았다”의 정확한 범위',2)
table(['층','관련 파일·기능','현재 상태'],[
('KG 원본과 그래프','kg_v2/authored, facts.jsonl, graph.json','종결 대응 보완까지 반영된 3,761진술. 최신 조언 검토에서는 수정하지 않음.'),
('기본 조회기','kg_v2/sentence_retrieval.py, query.py','기본형·형태·기능·조건·활용 관계 조회. 최신 조언 검토에서는 수정하지 않음.'),
('생성기','kg_v2/run_translation.py의 요청 구성·프롬프트·Qwen 호출','조회 자료를 전달하여 번역 생성. 최신 조언 검토에서는 수정하지 않음.'),
('실험 전용 확장','reports/kg_v2_factorial_20261012/','의미 검색과 변환 계획을 구현하고 200개 생성까지 완료. 기본 경로로 자동 채택하지 않음.'),
('검토·보완 제안','reports/kg_v2_advisory_20261012/','9개 제안과 반례·원문 대조·통계를 새로 작성. 모델 조언을 문헌 사실로 합치지 않음.')],[3.0,6.5,7.7])
p('따라서 “아무것도 수정한 적이 없다”는 뜻이 아니다. 이전 종결 대응 보완에서 KG와 조회기를 수정했고, 다음 단계에서는 별도 실험 코드를 추가했다. 이번 조언 검토의 생산 반영 여부를 말할 때만 “아직 변경하지 않았다”고 표현한다.')

h(chapters[2])
figure('pipeline',r'''digraph G {
graph [rankdir=TB,bgcolor="white",nodesep=.22,ranksep=.32];
node [shape=box,style="rounded,filled",fontname="Noto Sans CJK KR",fontsize=12,fillcolor="#EDF3F8",color="#4B6B84",margin=".16,.10"];
edge [fontname="Noto Sans CJK KR",fontsize=10,color="#5D7485"];
sources [label="구축 시점: 문헌 7종·474쪽\nAstra 작성·검토 → 코드 검증 → KG v2"];
input [label="표준어 발화"];
analysis [label="Qwen 분석 JSON + Kiwi 형태소\n원문 구간·시제·높임·화행"];
retrieve [label="기본: 형태·기능·관계 조회 + 조건 검사\n실험 B/D: 조건부 진술 의미 검색 추가"];
packet [label="A/B: 기존 참고 패킷\nC/D: 구간별 후보 + 보존 지침"];
generation [label="Qwen 방언 생성\n실제 요청·응답·조회 기록 저장"];
review [label="검토: 조언자 + 반례 검토자 + 문헌 대조자\n조정자 통합 → 아직 기본 실행 경로 미반영",fillcolor="#FFF3DD"];
sources -> retrieve [label="진술·근거"];
input -> analysis -> retrieve -> packet -> generation;
generation -> review [style=dashed,label="출력과 오류 추적"];
}''','그림 1. 현재 처리 흐름. 구축·검토용 Astra와 번역 런타임의 Qwen을 구분한다. B/D 및 C/D 확장은 별도 실험 경로다.')
table(['모델·도구','입력 → 출력','현재 호출 위치'],[
('Astra: 문헌 작성·검토','PDF 텍스트·표 이미지·작성 계약 → 진술 JSON·페이지 대장·검토 보고','지식 구축 시점. 새 문장을 번역할 때마다 호출하지 않음.'),
('Qwen3.8-27B: 분석','표준어 + 특징 정의/스키마 → clauses/source_units JSON','새 입력 분석. 최신 실험 기존38개는 캐시 공유, 추가12개는 새 분석.'),
('Kiwi','표준어 → 형태소·품사·문자 위치','로컬 형태 분석. Qwen 분석과 상호 보완하지만 오분석 가능.'),
('multilingual-e5-small','질의/지식 진술 텍스트 → 정규화 벡터','B/D 실험에서 로컬 CPU로 실행. 생성 모델이 아님.'),
('Python 조회·검사','분석·KG·조건 → 선택 후보와 제외 사유','구조화 조건 검사는 코드가 수행. 자유 서술 제약 전체를 자동 판정하지 않음.'),
('Qwen3.8-27B: 생성','표준어·분석·참고 패킷 → 방언 발화 문자열','실험에서 200개 성공, 202회 시도.'),
('외부 근거 없는 Astra high 비교','38개 입력을 한 작업 메시지로 전달 → 38개 번역','기록된 비교 조건. 도구·문헌 사용 금지. 순수 사전학습 base 모델이라는 뜻은 아님.'),
('사용자 제공 Sonnet 출력','사용자가 제공한 38개 번역','호출 설정·원프롬프트를 독립 확인하지 못함. 정답 자료로 간주하지 않음.'),
('세 역할 작업자 + 조정자','기존 출력·문헌·KG → 제안·반례·근거 검토','최신 검토는 세션 모델 상속, 별도 모델/추론 수준 재지정 없음.'),
('Neo4j','그래프 적재·조회·시각화','구조 확인용 별도 경로. 현재 문장 조회는 로컬 JSON을 Python으로 처리.')],[3.7,7.3,6.2])
p('현재 경로는 모델 가중치를 추가 학습하지 않는다. KG를 만드는 것, 임베딩 인덱스를 만드는 것, 프롬프트에 근거를 넣는 것은 각각 모델 학습과 다르다. 요청·응답의 모델명 일치는 확인 가능하지만 서버 내부 체크포인트의 불변성까지 입증하지는 못한다.')

h(chapters[3])
h('4.1 원문 범위와 현재 규모',2)
table(['문헌 / 자료','PDF 쪽','진술','주의할 범위'],[
(by_source[s['source_id']]['title'],s['pages'],s['facts'],
 '학술 자료. 지역·조사 시기·재인용 범위 확인' if by_source[s['source_id']]['source_type']!='web_document_snapshot' else '웹 문서의 로컬 PDF 스냅샷. 학술 검증과 동일하지 않음') for s in build['sources']],[9,1.5,1.5,5.2])
p('총 474쪽의 페이지 검토 대장이 있으며, 현재도 원문·자료 한계가 기록된 페이지가 있다. “전체 문헌을 작업 범위로 검토했다”와 “모든 적용 지식을 완전하게 추출했다”는 다르다. 최근에 가세예·보실래예 예문 누락을 발견한 사실이 이 차이를 보여준다. 페이지 번호는 인쇄 쪽수가 아니라 1부터 시작하는 PDF 쪽수다.')
type_labels={'context':'배경·문맥','usage':'용법','ending':'종결어미','phonology':'음운','construction':'구문','inflection':'활용','lexical_mapping':'어휘 대응','constraint':'제약 진술','particle':'조사'}
table(['진술 유형','수','의미'],[(k,v,type_labels[k]) for k,v in build['fact_types'].items()],[5.5,2,9.7])
p('진술 수는 출처·조건·표의 셀을 구별해 기록한 개수다. 3,761개를 서로 독립적인 변환 규칙 3,761개나 어휘 3,761개로 보고하지 않는다. 문헌의 출처 의존성도 있어 7종을 모두 독립 증거 7개로 계산하지 않는다.')
h('4.2 그래프가 연결하는 것',2)
node_counts=Counter(n['type'] for n in graph['nodes'])
node_meanings={'Source':'원문 문헌','SourcePage':'문헌의 PDF 페이지','KnowledgeStatement':'조건과 근거를 가진 지식 진술','ExpressionUse':'진술별 입력/목표 표현의 사용','Constraint':'자유 서술 제약','Evidence':'인용 근거','Function':'문법·화용 기능','Condition':'기계 판정 가능한 조건'}
table(['노드 유형','현재 수','역할'],[(k,v,node_meanings[k]) for k,v in node_counts.items()],[5.1,2.1,10])
figure('schema',r'''digraph G {
graph [rankdir=LR,bgcolor="white",nodesep=.18,ranksep=.35];
node [shape=box,style="rounded,filled",fontname="Noto Sans CJK KR",fontsize=11,fillcolor="#EDF3F8",margin=".13,.09"];
edge [fontname="Noto Sans CJK KR",fontsize=9,color="#607788"];
k [label="KnowledgeStatement\n지식 진술",fillcolor="#D7E9F2"];
x [label="ExpressionUse\n입력 표현 / 목표 표현"];
c [label="Condition / Constraint\n구조화 조건 / 서술 제약"];
f [label="Function\n문법·화용 기능"];
e [label="Evidence\n실제 인용문"];
p [label="SourcePage\nPDF 페이지"];
s [label="Source\n문헌"];
i [label="활용 진술\n예: 묵고"];
k -> x [label="source_expression\ntarget_expression"];
k -> c [label="requires / has_constraint"];
k -> f [label="describes_function"];
k -> e [label="supported_by"];
e -> p [label="located_on"];
p -> s [label="part_of"];
k -> i [label="candidate_has_inflection\n어휘 대응인 경우"];
}''','그림 2. 현재 graph.json의 관계 유형을 요약한 도식. 실제 노드 전체나 Neo4j 화면을 캡처한 그림이 아니다.')
p('어휘 대응의 목표 기본형과 활용 진술의 기본형이 맞으면 candidate_has_inflection으로 연결한다. 현재 이 관계는 484개이며 어휘 대응 44개에 연결되어 있다. 명사처럼 활용 대상이 아닌 항목도 있으므로 44/650을 동사 활용 포괄률이라고 부르면 안 된다. 경로가 연결되어 있어도 두 문헌의 뜻·지역·시대 범위가 일치하는지는 별도로 확인해야 한다.')
h('4.3 진술의 핵심 속성',2)
table(['속성','뜻과 주의점'],[
('id / source_id','진술과 문헌을 추적하는 식별자.'),('source_pattern / target_pattern','입력 쪽 표현과 방언 후보 또는 설명 대상. text·domain·lemma·sense·features를 구분한다.'),
('functions','평서·의문·요청·인용 같은 기능. 기능 일치만으로 자유 결합이 허용되는 것은 아니다.'),
('conditions','owner·field·op·value·basis로 구성된 검사 조건. 원문, 후보, 목표 말투, 문맥 중 무엇을 검사하는지 명시한다.'),
('constraints / scope','필요·배제·불확실성 등 설명과 지역·말씨·시대 범위. 대부분 자연어이므로 모든 내용이 코드로 검사되는 것은 아니다.'),
('evidence','원문 PDF 페이지와 인용 문자열. 문자열 존재와 진술 해석의 타당성은 별도 검증이다.'),
('claim_status','원문의 일반화, 실제 형태/예문, 저자 가설, 재인용, 분석자 해석, 부정 예문을 구분한다.'),
('usability','candidate / reference_only / blocked. candidate는 검색 후보이고 자동 생성 승인 상태가 아니다.'),
('notes','표준어 풀이의 작성 주체, 해석과 추가 한계 등을 남긴다.')],[5.2,12])
p('이 구조는 표현·조건·근거를 구분할 수 있다. 다만 모든 적용 관계가 충분히 구조화되어 있다는 뜻은 아니다. 표현 목록과 적용 규칙이 같은 후보 경로로 섞일 수 있고, 보존에 필요한 reference_only 지식은 생성 후보에서 제외될 수 있다. 현재는 이런 사용 경계부터 점검하는 것이 타당하다.')
citation('kg_v2/AUTHORING.md, fact.schema.json, build_summary.json, facts.jsonl, graph.json')

h(chapters[4])
h('5.1 Qwen 분석과 Kiwi가 각각 하는 일',2)
body('''Qwen 분석기는 원문을 번역하지 않고 JSON으로 분석한다. 절과 구성, 서술어, 평서/의문 등의 형식, 요청 같은 의도, 시제·양태, 청자 높임, 주체 높임을 기록한다. 원문에서 확인할 수 없는 정보는 추측해 확정하지 않도록 지시한다.

코드는 분석 문자열이 원문의 연속 부분인지 확인하고 문자 위치와 ID를 붙인다. Kiwi는 형태소·품사·위치를 따로 분석한다. 원문 분석은 가설이며, JSON 형식이나 문자열 결속 검증에 통과해도 의미적 분석이 정확하다는 뜻은 아니다.

최신 실험에서는 기존 38개의 저장 분석을 재사용하고 추가 12개를 새로 분석했다. 따라서 현재 분석 프롬프트가 요구하는 모든 구성 정보가 기존 캐시에 동일하게 들어 있다고 가정하면 안 된다. 또한 source_units의 자유로운 설명 전체가 v2의 구조화 조건으로 자동 연결되는 것은 아니다.''')
table(['분석 속성','설명','혼동하면 생기는 문제'],[
('host_form / host_pos','어미가 붙는 서술부의 기본형과 품사','문장 첫 동사를 최종 서술어로 선택하거나 파생하다를 놓칠 수 있음.'),
('sentence_type / speech_act','문장 형식과 발화 의도','질문 형식의 부탁을 평서/명령 하나로 축소할 수 있음.'),
('observed_politeness','원문 청자 높임','계시다만 보고 상대에게도 존댓말이라고 잘못 판단할 수 있음.'),
('observed_subject_honorific','행위 주체를 높이는 표현','청자 높임만 맞추면서 -시-나 계시다를 손실할 수 있음.'),
('observed_prefinals','원문에서 확인한 선어말어미 순서','사건이 과거라는 이유로 실제로 없는 -었-을 만들 수 있음.'),
('tense_or_modality','시제·양태 해석','시제의 의미와 표면 형태를 같은 것으로 처리할 수 있음.'),
('span / occurrence','원문 범위 / 동일 문자열 등장 순서','다른 절에 지식을 적용하거나 반복 표현 위치를 혼동할 수 있음.'),
('source_units / parent_index','관형·인용·보조용언 등 구성과 포함 관계','부분 어미만 치환하여 상위 구문의 의미를 손상할 수 있음.')],[4.5,6.1,6.6])
h('5.2 기본 조회기의 순서',2)
body('''① 원문에 결속한 어휘·서술부를 찾는다. 어휘는 기본형 대응을, 문법은 형태와 기능을 주요 단서로 사용한다.

② 어휘 대응을 찾은 뒤 graph 조건에서는 candidate_has_inflection 관계를 통해 실제 활용 진술을 추가 조회한다. 표에 없는 모든 활용을 자유 생성하는 문법 엔진은 아니다.

③ candidate 여부, 부정 예문 여부, 구조화 조건의 충족·실패·미확인을 검사한다. 조건 검사 함수는 Python이며, Qwen이 이 코드를 대신 실행하는 구조가 아니다.

④ 조회 경로의 우선순위와 원문의 말투에 따른 순위를 적용한다. 구체적인 종결 대응을 우선하고 동일 구간의 넓은 기능 후보를 억제하는 규칙, 일부 경로의 절별 분산, 중복 표현 억제도 있다.

⑤ 같은 구간 최대 6개와 전체 참고 패킷 24,000자 제한 안에서 후보를 선택하고 제외 사유를 기록한다. 경로용 2,000자를 유보한다.

⑥ 후보와 적용 구간·검사 결과·출처를 묶어 Qwen에 전달한다. 자유 서술 제약과 의미·지역 적합성의 상당 부분은 여전히 생성기의 해석에 남는다.''')
h('5.3 조건 검사의 범위',2)
table(['구분','현재 처리'],[
('검사 대상 owner','input=원문 특징, candidate=선택한 방언 어휘·어간, target=목표 말투, context=상황.'),
('지원 연산','eq(같음), in(집합 안), not_in(집합 밖), exists(값 존재 여부).'),
('정보 미확인','unknown으로 남기며 문장 생성 참조 선택 단계에서는 미확인 구조화 조건 후보를 제외한다.'),
('자유 서술 제약','검토 필요 표시를 남긴다. “자연스러움”, 문맥별 뜻과 지역 선호를 모두 판정하는 코드가 없다.'),
('빈 conditions','검사할 명시 조건이 없다는 뜻. 모든 입력에 적용 가능하다는 증명은 아니다.'),
('reference_only','현재 직접 변환 후보에서 제외된다. 이것이 관형형 보존 지식 전달에는 장애가 될 수 있다.')],[4.6,12.6])
p('현재 파생 서술어 누락의 구체적 사례는 출발/NNG+하/XSV, 시작/NNG+하/XSV다. 조회용 서술어 품사 목록에 XSV가 없어 DEV011과 NEW004는 A/B 모두 후보가 없다. 같은 분석 경로를 공유하는 의미 검색만 추가해도 이 누락은 복원되지 않는다.')
citation('kg_v2/sentence_retrieval.py, query.py; current_kg/experiments/run_auto_analysis_development.py; factorial review_metrics.json')

h(chapters[5])
h('6.1 전체 문장의 주제 대신 변환 기능을 검색',2)
body('''예를 들어 “아까 문 앞에 있던 사람이 제 친구예요.”의 어미를 바꿀 때 친구·문 앞이라는 주제보다 명사 서술·현재·평서·청자 높임·예요라는 정보가 중요하다. 이번 실험은 어휘 질의와 문법 질의를 나누어 만들었다. 생성할 때는 여전히 전체 원문을 함께 제공한다.

검색 문서는 개별 엣지 한 줄이 아니라 지식 진술 중심이다. 입력/목표 표현, 문법 기능, 적용 조건, 사용 범위, 추가 제약을 텍스트로 구성해 임베딩했다. 원문 인용 본문은 벡터화하지 않았고, 근거 페이지는 인덱스 기록에, 전체 인용은 KG와 조회 로그에 보존했다.''')
table(['속성','실제 설정'],[
('임베딩 모델','intfloat/multilingual-e5-small, 로컬 CPU'),
('고정 revision','614241f622f53c4eeff9890bdc4f31cfecc418b3'),
('인덱스','3,761개 진술 / 3,783개 청크 / 여러 청크인 진술 22개'),
('분할','진술 텍스트를 384토큰 단위로 분할. 최장 진술 621토큰. 인용문 전체를 추가로 붙이지 않음.'),
('벡터 계산','passage: / query: 접두어. 마스크를 적용한 평균 풀링 후 정규화. 여러 청크는 합산 후 정규화하여 진술 하나로 표현.'),
('유사도','코사인 유사도. 같은 정규화 벡터 공간에서 질의와 진술 비교.'),
('추가 발견 후보','허용 풀 안의 구간별 상위 12개를 추가 검토. 최종 전달 상한 6개와 다른 숫자.'),
('기존 제한 유지','어휘 기본형 일치, 문법 기능·말투 필터, 구조화 조건 검사 유지.'),
('최종 순위','기존 경로·말투 우선순위 안에서 의미 유사도를 반영. 전체 KG를 유사도 하나로 정렬한 실험이 아님.')],[4.1,13.1])
p('선행 KG 질의응답 연구에서 착안한 것은 그래프 정보를 읽기 쉬운 텍스트로 표현하고 의미적으로 검색한다는 방향이다. 이번 실험은 그 논문의 KGQA 태스크나 모든 알고리즘을 재현한 것이 아니다. 문헌의 번역 적용 조건을 다루는 별도 설계로 구체화했다.')
h('6.2 C/D의 구간별 변환 계획',2)
body('''A/B의 references는 후보마다 근거 인용과 검사 로그까지 포함한다. C/D는 같은 span의 후보를 하나의 그룹으로 묶고, 원문 구간과 시제·청자 높임·주체 높임·기능 보존값을 그룹에 붙였다.

A와 C, B와 D 사이에서 선택된 진술 ID와 적용 구간, 표현·조건·제약·범위·주석·관계 경로는 유지했다. 원문 인용 본문과 행정적 검사 로그는 생성 입력에서 생략했고 출처 ID와 페이지는 남겼다. 따라서 형식만 달라지는 완전한 무손실 실험이나 순수한 토큰 길이 실험은 아니다.

이 계획은 조건부 후보 목록이지, 모든 대상 구간에 대해 정확한 최종 활용형 하나를 제공하는 정답 계획은 아니다. 기존 종결 대응에서 구현한 제한적인 조합만 사용하고 문헌의 실제 완성형과 규칙 조합형을 구분한다.''')
h('6.3 DEV004 실제 전달 사례',2)
p('원문: '+input_by['DEV004'])
dev4=next(j for j in jobs if j['id']=='DEV004' and j['condition']=='C')
plan=json.loads(dev4['request']['messages'][1]['content'].split('\n',1)[1])['transformation_plan'][0]
table(['실제 속성','저장된 값 또는 설명'],[
('target_text / span',str(plan['target_text'])+' / '+str(plan['span'])),
('후보 수',len(plan['candidates'])),
('후보 표현',' / '.join(x['target_pattern']['text'] for x in plan['candidates'])),
('보존값','명사 서술·평서·현재·청자 높임. 주체 높임은 not_observed.'),
('구간 범위 주의','어미 정보가 들어 있어도 이 항목의 span 자체는 문장 전체 [0,22)다. 모든 후보 구간이 최소 어미 길이로 잘린 것은 아니다.'),
('후보 배열 길이','A/B 5,119자, C/D 3,679자. 출력은 네 조건 모두 같은 친구라예 문장.')],[4.3,12.9])
p('첫 후보의 ending_conversion은 replace_copula_and_ending 조합이다. derived_composition_not_attested_surface는 이 입력의 완성형을 문헌에서 그대로 인용했다는 뜻이 아니라 기존 대응을 적용해 조합한 것임을 표시한다.')
citation('factorial semantic_retrieval.py, experiment.py, semantic_index.json, jobs.json')

h(chapters[6])
table(['조건','조회 방식','생성 전달 형식','통제 관계'],[
('A','기존 KG v2 그래프 조회','기존 references + graph_paths','현재 기준선'),
('B','A에 조건부 진술 의미 검색 추가','기존 references + graph_paths','A와 조회 방식 비교'),
('C','A와 같은 선택 후보','구간별 transformation_plan + graph_paths','A와 전달 방식 비교'),
('D','B와 같은 선택 후보','구간별 transformation_plan + graph_paths','B와 전달 방식, C와 조회 방식 비교')],[1,5,6.2,5])
table(['실행 항목','확인된 값'],[
('입력','DEV30 + LONG8 + NEW12 = 50개 입력. 200개의 서로 다른 입력이 아님.'),
('NEW12의 지위','출력 확인 전에 고정했지만 기존 실패 유형을 참고해 만든 도전 입력. 독립적으로 수집한 일반화 평가셋이 아님.'),
('분석','기존38개 캐시 공유 + 새12개 분석. 새 분석 호출 총15회.'),
('분석 재시도','NEW010 필수 필드 누락으로 총3회, NEW012 허용되지 않는 라벨로 총2회. 형식 오류만 피드백.'),
('생성','200개 최종 성공, 202회 시도. DEV016 A/D 시간 초과 후 각1회 재시도. 품질을 보고 골라 재생성하지 않음.'),
('생성 설정','Qwen3.8-27B / seed 20260923 / temperature 0.3 / top_p 0.8 / top_k 20 / max_tokens 512 / enable_thinking false'),
('반복','조건당 입력별1회. 반복 실험 또는 통계적 유의성 검증이 아님.'),
('고정한 자료','KG·분석·기본 코드·원문·요청 해시. A/C 및 B/D 후보 일치 검증.'),
('이전 결과와 관계','기존38개의 A 요청은 직전 ending_repair v2_graph 요청과 같으며 이번에도 새 생성. 해당38개 출력은 직전 출력과 동일.'),
('평가 지위','오류 추적용 개발 비교. 독립 화자 정답률이나 공인 벤치마크 점수 없음.')],[4.2,13])
h('7.1 모델 간 비교를 해석하는 방법',2)
body('''Astra high와 사용자 제공 Sonnet 결과는 어떤 부분을 더 자연스럽게 표현할 수 있는지 찾는 비교 자료로 유용하다. 그러나 Qwen+KG와 차이가 난다고 KG만을 원인으로 결론 내릴 수 없다. 모델 자체, 추론 설정, 요청 단위와 프롬프트가 함께 다르기 때문이다.

Astra 비교는 외부 도구·문헌 없이 38개를 한 메시지로 처리하도록 요청했다. Qwen은 입력별로 생성했고 thinking을 껐다. Sonnet은 사용자가 제공한 출력이며 원호출 설정을 검증하지 못했다. 외부 근거를 사용하지 않았다는 뜻과 순수 사전학습 base 체크포인트라는 뜻은 다르다.

강한 모델의 출력도 의미 보존 측면에서 점검해야 한다. 사용자 Sonnet DEV002의 집에 오는 길은 원문 집에 가는 길과 이동 관점이 다르고, DEV006의 할매요는 할머니께서라는 주어를 호격으로 바꾼다. 방언성이 두드러진다는 이유로 전체를 정답으로 채택하지 않는다.''')
citation('factorial protocol.json; astra_high_no_kg protocol.json; model_comparison comparison.json')

h(chapters[7])
h('8.1 측정 대상과 단위',2)
body('''저장된 jobs.json의 실제 user 메시지를 역직렬화하여 후보를 집계했다. A/B 패킷은 retrieval.json과 일치하는지 확인했다. 임의 후보나 추정 토큰수는 사용하지 않았다. 문자수는 ensure_ascii=False, separators=(",", ":")로 직렬화된 실제 JSON의 Python len 값이다.

후보를 받은 고유 span만 전달 구간으로 센다. 어휘 구간이 문장 전체 구간 안에 포함되어도 서로 다른 span이면 두 구간이다. 원문을 서로 겹치지 않게 나눈 조각 수나 모든 잠재 검색 단위 수가 아니다. 같은 진술이 다른 구간에 전달되면 각각 후보 1개로 센다.

50개 입력에는 여러 문장으로 된 LONG과 일부 NEW가 포함된다. 따라서 “표준어 문법 문장 한 개당”이라고 바꾸어 부르면 안 된다. 아래 수치는 실험 입력 ID 하나당 통계다. 후보가 전혀 없는 DEV011·DEV018·NEW004도 입력당 평균 계산의 분모에 포함했다.''')
table(['지표','정의 / 계산','읽는 방법'],[
('입력당 평균 구간 수','해당 조건 전체 전달 구간 수 ÷ 입력 수','A: 150 ÷ 50 = 3.00개. 0구간 입력도 포함.'),
('구간당 평균 후보 수','전달 후보 수 ÷ 후보가 있는 전달 구간 수','A: 412 ÷ 150 = 2.7467개. 입력별 평균을 다시 평균한 값과는 다름.'),
('구간 후보 묶음 평균 길이','각 구간의 후보 배열 JSON 길이를 구한 뒤 구간 전체에서 평균','출처·인용·키·구두점·이스케이프 포함. 공유 그래프 경로와 전체 분석은 제외.'),
('입력별 최대 묶음의 평균','입력마다 가장 긴 구간 후보 배열 길이를 구한 뒤 입력 전체에서 평균','“평균 최대”를 이 의미로 정의. 후보 없는 입력의 최대는0자.'),
('관측 최대 구간 길이','해당 조건 모든 입력·구간 중 가장 긴 후보 배열','실제로 관측한 최대값이며 설정 한도가 아님.'),
('구간 표현 전체 길이','C/D에서는 후보 배열에 target_text·span·preserve 등 그룹 공통정보를 포함','후보 배열 길이보다 길다. A/B는 후보 배열 길이와 같음.'),
('참고 패킷 길이','A/B references+graph_paths, C/D transformation_plan+graph_paths+plan_policy','구간 여러 개와 공통 경로를 함께 포함.'),
('전체 user 메시지 길이','지시문 + 원문 + 분석 + 형태소 + 참고 패킷','system 메시지는 제외. 모델 전체 컨텍스트나 입력 토큰수와 다름.'),
('구간별 가장 긴 단일 후보의 평균','구간마다 후보 객체 중 가장 긴 하나를 찾고 전체 구간에서 평균','입력별 최대 “묶음” 평균과 구별하기 위한 보조 지표.')],[4.2,7,6])
h('8.2 전체 50개 입력: 후보 수',2)
s50=stats['summary']['all50']
table(['조건','입력 수','전달 구간 총수','입력당 구간','후보 총수','구간당 후보'],[
(c,r['input_count'],r['delivered_segment_count_total'],f"{r['delivered_segments_per_input_mean_including_zero']:.4f}",r['candidate_count_total'],f"{r['candidates_per_delivered_segment_pooled_mean']:.4f}") for c,r in s50.items()],[1.4,2,3.4,3.4,3.4,3.6])
h('8.3 전체 50개 입력: 구간별 문자 길이',2)
table(['조건','후보 묶음 평균','입력별 최대 묶음의 평균','관측 최대 묶음','공통정보 포함 평균'],[
(c,f"{r['candidate_array_characters_per_segment_mean']:,.2f}",f"{r['per_input_largest_segment_characters_mean_including_zero']:,.2f}",f"{r['candidate_array_characters_largest_segment']:,}",f"{r['segment_representation_characters_mean']:,.2f}") for c,r in s50.items()],[1.2,4,4,4,4])
p('단위는 모두 문자(자). C/D의 후보 배열은 공통 보존 지침을 제외하므로 그 자체만을 완전한 구간 입력 분량이라고 부르지 않는다. 공통정보를 포함하면 C는 평균3,206.79자, D는3,221.51자다. 모든 조건에서 최대 후보 묶음이 나온 입력은 DEV024의 전체 절 구간이다.')
table(['조건','구간별 가장 긴 단일 후보의 평균'],[(c,f"{r['per_segment_largest_single_candidate_characters_mean']:,.2f}자") for c,r in s50.items()],[3,14.2])
h('8.4 원래 38개와 추가 12개를 분리한 통계',2)
for cohort,label in [('previous38','기존 DEV30+LONG8: 38개'),('new12','추가 NEW12: 12개')]:
    p(label)
    table(['조건','입력당 구간','구간당 후보','묶음 평균자','최대 묶음 평균자','관측 최대자'],[
    (c,f"{r['delivered_segments_per_input_mean_including_zero']:.4f}",f"{r['candidates_per_delivered_segment_pooled_mean']:.4f}",f"{r['candidate_array_characters_per_segment_mean']:,.2f}",f"{r['per_input_largest_segment_characters_mean_including_zero']:,.2f}",f"{r['candidate_array_characters_largest_segment']:,}") for c,r in stats['summary'][cohort].items()],[1.2,2.8,2.8,3.3,3.8,3.3])
p('두 집합의 차이는 입력 구성의 차이를 포함한다. 추가12개가 별도 모델 설정으로 더 많은 지식을 받았다는 뜻은 아니다. 전체50개 평균만 보거나 원래38개 평균만 볼 때 분모를 반드시 함께 명시한다.')
h('8.5 입력 전체 길이와 24,000자 제한',2)
table(['조건','참고 패킷 평균','참고 패킷 최대','user 메시지 평균','user 메시지 최대'],[
(c,f"{r['packet_characters_mean']:,.2f}",f"{r['packet_characters_max']:,}",f"{r['user_message_characters_mean']:,.2f}",f"{r['user_message_characters_max']:,}") for c,r in s50.items()],[1.2,4,4,4,4])
body('''24,000자는 구간별 제한이 아니라 입력 전체 references+graph_paths에 적용하는 예산이다. 후보 선택 중에는 관계 경로를 위해 2,000자를 유보한다. C/D는 이 선택 결과를 계획 형식으로 다시 직렬화하며, 그 재구성 패킷에 별도 24,000자 제한을 다시 적용하지 않는다.

원문 분석과 지시문이 추가되므로 전체 user 메시지는 24,000자를 넘을 수 있다. 이 현상을 예산 검사 실패라고 해석하면 안 된다. 문자와 토큰도 다르므로 여기서 토큰 비용이나 모델 컨텍스트 사용률을 역산하지 않는다.

A→C, B→D는 후보 수가 같아도 전달 길이가 줄어든다. 이는 계획 형식에서 원문 인용 본문·검사 로그를 덜 보내기 때문이다. 길이 효과, 정보 생략 효과, 보존 지침 추가 효과가 함께 바뀌므로 하나의 원인으로 분리해 해석할 수 없다.''')
h('8.6 후보 상한과 길이 제한은 얼마나 작동했는가',2)
table(['항목','A/C 선택 기반','B/D 선택 기반'],[
('실제6개 전달 구간','35 / 150 = 23.33%','37 / 160 = 23.13%'),
('개수 제한 탈락 기록이 있는 입력','31 / 50','33 / 50'),
('패킷 길이 제한 탈락 기록이 있는 입력','10 / 50','12 / 50'),
('후보0개 입력','DEV011, DEV018, NEW004','동일')],[7,5.1,5.1])
table(['구간의 후보 수','A/C 구간 수','B/D 구간 수'],[(n,s50['A']['candidate_count_distribution'].get(str(n),0),s50['B']['candidate_count_distribution'].get(str(n),0)) for n in range(1,7)],[7,5.1,5.1])
p('평균이6보다 작아도 약23%의 구간은 상한에 도달했다. 다만 탈락 기록은 고유한 유효 규칙의 수나 꼭 필요한 후보를 놓친 횟수가 아니다. 동일·유사 후보와 문맥 부적합 후보도 포함될 수 있다. 상한 변경의 효과를 판단하려면 탈락 후보의 필요성을 표시하고 고정된 후보 풀에서 별도 실험해야 한다.')
citation('advisory measure_retrieval.py, retrieval_statistics.json, segment_lengths.tsv, input_lengths.tsv. 실제 jobs.json 기준 집계.')

h(chapters[8])
h('9.1 전체 비교에서 확인할 수 있는 것',2)
table(['비교 / 관찰','실제 수','해석'],[
('A↔B 후보 집합 변화','37 / 50 입력','검색 결과가 달라졌다는 증거. 회수율이나 번역 개선율이 아님.'),
('A↔B 출력 문자열 차이','20 / 50 (같음30)','벡터 검색으로 다른 후보가 들어가도 같은 출력 가능.'),
('C↔D 출력 문자열 차이','15 / 50 (같음35)','계획 형식에서도 검색 변화가 출력 변화로 항상 이어지지 않음.'),
('A↔C 출력 문자열 차이','25 / 50','동일 후보의 전달 형식 변화와 함께 출력이 달라진 경우.'),
('B↔D 출력 문자열 차이','20 / 50','형식·보존 지침·인용 생략을 묶은 효과. 순수 길이 효과 아님.')],[6,3.8,7.4])
body('''A/B가 비슷하고 C/D가 비슷한 것은 동일한 Qwen과 원문 분석을 공유하고, B/D에서도 기존 경로 우선순위와 적용 필터를 유지했기 때문일 수 있다. 검색 후보가 바뀌어도 생성기가 같은 대안을 고르거나 후보를 사용하지 않을 수 있다. 이 설명은 기록에 부합하는 해석이며 각 원인의 기여율을 인과적으로 분리한 결과는 아니다.

표준어와 같은 표현이 남는 현상을 모두 실패로 세지 않는다. 문헌에도 해예체 사용자가 해요체를 함께 쓴다고 설명한다. 다만 변환 목표에 필요한 후보를 반복해서 놓치거나 높임·의미를 손상하는 경우는 별도 문제다.''')
output_case('DEV004','네 조건 모두 명사 서술의 청자 높임을 유지한 친구라예를 생성했다. 실제 종결 대응과 조합 계획이 전달되는 사례다. 이 한 사례로 일반 종결어미 문제가 해결됐다고 볼 수는 없다.')
output_case('DEV011','네 조건 모두 청자 높임이 약화된다. A/B 후보가0개이며 파생 하/XSV의 조회용 서술어 누락과 연결된다. 조회되지 않은 단계와 생성 모델의 높임 선택을 분리해 진단해야 한다.')
output_case('DEV006','B/C/D는 계세요로 청자 높임을 유지한다. 시고의 지역 적합성·자연스러움까지 검증된 것은 아니다. 청자 높임 개선과 전체 문장 품질 개선은 구분한다.')
output_case('DEV008','C에는 강우 서술과 문장 구조의 손상이 보인다. 짧게 정리한 계획 형식이 항상 더 나은 결과를 보장하지 않는 사례다.')
output_case('NEW005','C/D의 될까예는 문의 표현을 보존하는 국소적 이점이 보인다. 지역·화자에 따른 자연스러움의 독립 평가까지 끝난 것은 아니다.')
output_case('LONG006','A/B는 관형형을 종결형처럼 바꾸어 가노 길·좋아하노 빵을 만든다. C/D는 관형 구조를 회복하지만 원문의 가는을 가던으로 바꾼 시제·상 차이가 남는다.')
output_case('LONG008','B/D에는 -지예 후보가 전달되지만 첫 질문은 여전히 힘드셨제/힘드셨죠다. D의 말씀드릴래예는 목록의 -래예와 약속 표현의 적용 조건 간 간격을 보여준다. 특정 진술 하나가 원인임을 확정하는 제거 실험은 아직 하지 않았다. 앞에라는 어휘만으로 이 문장을 확정 오류 처리하지 않는다.')
h('9.2 반드시 반영할 기존 평가 정정',2)
table(['기존 판단','원문 확인','현재의 올바른 해석'],[
('앞에 갈게요는 먼저의 의미를 공간 방향으로 잘못 바꿨다.','RYU2014 PDF27에 먼저→앞에 및 앞에 가요가 있음.','시간적 용법이 있으므로 자동 오류로 단정하지 않음. 문맥·지역·세대 적합성은 추가 검토.'),
('맛있대는 직접 경험을 전해 들은 말로 바꾼다.','문법 PDF28의 -더라 대응에 -대/-더라/-더래가 있음. 예문은 덥데.','회상 대응 가능성을 인정. 표기 대/데 차이와 해당 출력의 적합성은 검토 필요.'),
('현재도 먹고→묵고 경로가 없다.','LONG001에 묵고 활용 진술이 실제 전달됨.','이 사례는 조회 성공 후 생성 적용을 확인. 다른 활용까지 모두 해결된 것은 아님.'),
('맛있지예 표현 자체가 KG에 없다.','문법 어휘 근거와 높임 PDF4 review0099에 존재.','인용문 존재와 해당 구성에 쓸 독립 적용 지식 연결을 구별.')],[5.2,6,6])
p('원래 출력과 로그를 수정하지 않고 해석 정정을 별도 기록했다. 과거 분석의 확정 오류 라벨을 그대로 합산하면 잘못된 품질 비교가 된다. 이번 보고서는 정정된 해석을 따른다.')

h(chapters[9])
table(['역할','실제 수행','결과'],[
('변환 조언자','Astra·Sonnet·Qwen 출력과 현재 실패 유형을 비교하여 적용 조건과 보존 지침 제안','제안9건. 모두 모델 조언으로 표시.'),
('반례 검토자','규칙의 과잉 일반화·문맥 충돌을 검토하고 최소 대조 입력 구성','문제군10개, 새로운 입력20개, 제안9건 검토.'),
('문헌 대조자','원문·KG·실제 조회 기록을 비교하고 PDF 표·예문 확인','쟁점10개, 인용22개, 실제PDF7종 해시 대조.'),
('조정자','ID·인용·누락 예문 검색·원본 불변성을 검증하고 교차 지적 반영','reviewed_advice.json, coordinator_validation.json, 통계와 정정 보고서.')],[3.2,8,6])
p('이 역할 분담은 모델의 사전학습 지식을 활용하되 그것을 문헌 근거로 위장하지 않기 위한 절차다. 같은 계열 세션 모델의 역할 분리는 사람 전문가의 독립 검증이 아니다. 최신 검토 작업자에게 별도의 Astra high 설정을 지정한 것으로 보고하지 않는다.')
h('10.1 제안 9개의 요약과 상태',2)
adv_titles=['파생·복합 서술어 복원','의문 유형과 높임 문의','세요의 요청·서술 구별','죠/네요의 화용 기능','목록과 실제 결합 조건 분리','높임·겸양·주어/호격 보존','관형절 보존','어휘 의미 역할과 적용 범위','인용·경험·전언 구분']
review_map={r['proposal_id']:r for r in critic['advisor_review']}
table(['ID / 주제','핵심 제안','반례 검토 상태'],[
(a['id']+'\n'+name,a['proposed_rule'],'설계 방향 수용·진단 대상' if review_map[a['id']]['decision']=='accept' else '조건 보완 후 진단') for a,name in zip(advisor['proposals'],adv_titles)],[3.9,10,3.3])
p('수용3건은 분석·보존 설계 방향을 수용한 것이다. 모델이 만든 방언 예문을 문헌의 확인형이나 사람 정답으로 승인한 것이 아니다. 나머지6건은 조건을 좁히는 작업이 더 필요하다. 모든 항목의 production_KG_active는 false다.')
h('10.2 문헌 대조의 중요한 발견',2)
table(['쟁점','현재 확인','수정 대상'],[
('가세예 / 보실래예','높임 PDF26에 실제 존재. 모든fact 문자열 공백 정규화 검색에서0건.','원문 사례 추가 추출 후보. 동등 추상 규칙 전체의 부재까지 증명하지 않음.'),
('맛있지예','문법 full0941~0943, 높임 review0099에 존재. PDF26의 별도 적용 연결은 미확인.','기존 근거와 해당 문법 구성을 연결할 필요.'),
('관형형 보존','RYU full0098에 정보가 있으나 LONG006 조회에 빠짐.','변환 후보와 별도로 보존 지식을 전달하는 통로.'),
('까요','관련 문헌 형태가 있는데 DEV013 문법 후보0개.','의문 구조·간접 요청·형태 정규화 연결.'),
('래예','명사서술 모양이래예와 의향 질문 보실래예는 서로 다른 구성.','같은 표면 문자열만으로 목록을 임의 활용 규칙으로 쓰지 않기.'),
('먹다→묵고','LONG001에서 실제 활용 근거를 전달.','기존 성공 경로를 회귀 기준으로 보존하고 생성 적용 단계 점검.')],[3.2,7.5,6.5])
p('문헌 대조 범위는 관련 쟁점의 원문·표 확인이며, 이번에474쪽을 다시 의미 전수 검수한 것은 아니다. 조정자의22개 인용 문자 일치 검사는 원문 존재 확인이다. 적용 조건의 모든 일반화가 옳다는 보증과 구분한다.')
h('10.3 조언을 어떻게 보완 지식으로 관리하는가',2)
body('''문헌에서 직접 확인한 내용은 원문 인용과 페이지를 가진 진술로 관리한다. 문헌 규칙을 이 입력에 적용해 만든 형태는 조합형이라고 기록한다. 문헌에서 확인되지 않은 모델의 규칙·예문은 model_advice_not_documentary_fact로 유지한다.

원문이 충분하지 않은 부분도 조언자로부터 적용 조건과 반례를 갖춘 가설을 받을 수 있다. 이 때문에 모델 조언을 쓰지 못하는 것은 아니다. 다만 문헌 기반 조건과 모델 보완 조건의 실험을 구분해야 어느 지식이 효과를 냈는지 설명할 수 있다.

현재 reviewed_advice.json은 제안·반례 검토·관련 문헌 쟁점·조정 상태를 함께 묶은 보완 자료다. 기본 KG의 자동 생성 후보에는 아직 편입하지 않았다. 새20개 입력에도 방언 정답을 미리 부여하지 않았고 생성은 실행 전이다.''')

h(chapters[10])
table(['원인 후보','증거 수준','현재 근거','아직 필요한 확인'],[
('분석→조회 연결 누락','구체 사례 확인','출발/시작의 XSV 호스트 누락과 후보0개.','복원 후 필요한 후보 회수·과거/높임 보존 회귀.'),
('원문 사례의 KG 누락','특정 문자열 사례 확인','가세예·보실래예 원문 존재, fact 검색0건.','예문·상황·결합 조건을 함께 추출하고 추상 규칙 중복 확인.'),
('보존 지식의 전달 누락','구체 사례 확인','reference_only 관형형 보존 정보가 LONG006에 선택 안 됨.','별도 보존 경로가 관형절 손상을 줄이는지 비교.'),
('관련성과 적용 가능성 혼동','강한 진단 근거','래예 목록과 약속 발화 적용 간 간격.','해당 후보 제거/조건 강화 실험으로 인과 기여 확인.'),
('자료 분량과 후보 제한','제한 작동 확인, 품질 인과 미확인','6개 상한 구간 약23%, 길이 탈락 입력 A10/B12.','필요 후보 누락을 표시하고 후보수·길이 요인을 분리.'),
('KG 노드·관계 설계 불충분','일부 표현 경계의 개선 필요','목록/적용규칙 구분과 보존용 연결이 부족.','전체 재설계보다 진술 유형·조건·조회 채널 수정의 효과 먼저 확인.'),
('생성 모델의 지식 사용','관련 관찰 있음','지예 노출 후에도 미사용, 같은 후보의 형식에 따른 출력 변화.','정확한 간결 계획을 제공한 통제 실험 후 모델·추론 조건 비교.'),
('평가자의 방언 해석 오류','실제 정정 완료','앞에·대의 문헌 용법을 놓쳐 확정 오류 처리.','문헌 범위와 화자 평가를 함께 반영한 판정 기준.')],[3.4,3.2,5.2,5.4])
h('11.1 구조화 조건이 빈 경우의 해석',2)
table(['실제 전달 참조','A','B'],[
('전체',412,439),('conditions가 빈 참조',347,362),('자유 서술 constraints가 있는 참조',410,434)],[9.2,4,4])
p('이는 KG 전체의 비율이 아니라 해당 실험에서 구간에 결속해 전달된 참조의 수다. 어휘 대응 중에는 추가 조건 없이 쓸 수 있는 것도 있으므로 빈 조건을 모두 잘못된 지식으로 세면 안 된다. 다만 “조건 검사 통과”가 실제 문법·화용 적합성까지 보장한다고 보기 어렵다는 점을 보여준다.')
h('11.2 확정할 수 없는 결론',2)
p('현재 증거만으로 문헌이 부족한 것이 유일한 원인, KG 구조가 전부 잘못된 것이 유일한 원인, Qwen의 능력이 유일한 원인이라고 말할 수 없다. 단계별 누락과 잘못된 적용, 생성 선택, 평가 해석이 함께 작동한다. 각 수정이 어느 단계에 영향을 주는지 기록해야 다음 결과를 설명할 수 있다.')

h(chapters[11])
p('아래는 후속 구현·실험안이다. 이 보고서를 작성하면서 새 번역을 실행하거나 생산 KG·조회기를 수정하지 않았다.')
table(['순서','수정 내용','기존 KG 수정 여부','통과 기준'],[
('1','오류 판정 정정과 현재 기준선 고정','없음','앞에·대를 일괄 오류 처리하지 않고 기존 출력/해시 유지.'),
('2','파생·복합 서술어 분석과 조회 연결 복원','주로 분석·조회 코드','출발/시작 등 검색 대상 생성, 시제·높임 원문 결속 유지.'),
('3','reference_only 보존 지식의 별도 전달','주로 조회·패킷; 근거 연결 보완 가능','관형형 보존 정보가 필요한 구간에 전달되고 치환 후보와 구별됨.'),
('4','가세예·보실래예 등 확인된 원문 사례 추가','버전 기록 후 KG 보완','원문 페이지·조건·주체/청자 높임·질문 기능을 함께 기록.'),
('5','목록/예문/규칙의 적용 조건과 우선순위 개선','지식 표현·조회 모두 가능','같은 래예의 다른 구성 구분, 불확실한 일반화 보류.'),
('6','구간별 계획과 생성 후 보존 검사','생성 입력·검사 코드','높임·부정·시제·이동 방향·인용·관형 역할 오류를 구간별 표시.'),
('7','후보수/길이·모델 보완의 효과 분리','실험 설정과 별도 조언 자료','변경 축을 통제하고 회수·적용·보존·자연스러움 지표 분리.')],[1,5.9,4.2,6.1])
h('12.1 후보 개수와 길이를 분리하는 실험안',2)
body('''후보 개수 실험에서는 원래의 후보 풀과 순위·생성 설정을 고정하고 같은 구간의 상한만 바꾼다. 예를 들어3/6/9개를 탐색할 수 있으나 이 값이 최적이라는 근거는 아직 없다. 전달하지 않은 후보가 어떤 입력의 어떤 기능에 필요한지를 사전에 표시한다.

길이 실험에서는 선택 후보 ID와 필수 조건을 고정하고, 인용과 설명의 표시 수준만 바꾼다. 조건 삭제와 길이 단축이 동시에 일어나지 않도록 보존할 필드를 정한다. 현재 C/D는 인용 생략과 보존 지침 추가까지 함께 바뀌므로 이 실험을 대신하지 못한다.

모델 조언 실험에서는 문헌만 사용하는 조건과 문헌+검토된 모델 제안을 사용하는 조건을 구분한다. 모델 제안 출처를 남기면 부족한 문헌 지식 보완의 효과와 원래 KG 조회 개선의 효과를 분리할 수 있다. 특정 모델 출력 문장을 그대로 외워 평가 입력에 넣지 않는다.''')
h('12.2 생성 후 보존 검사',2)
p('현재는 사후 오류 검토와 일부 분석/형태 검사만 있다. 종합적인 의미 보존 검사와 자동 국소 재생성은 아직 구현 완료 상태가 아니다. 후속 검사는 원문의 주어·호격, 주체/청자 높임, 과거·부정 범위, 질문/요청/약속, 이동 방향, 인용 내용과 보고 시제를 분리해 확인해야 한다. 불확실한 자동 표시는 사람 검토 대상으로 남기며 문자열 하나만 보고 합격·실패를 확정하지 않는다.')
table(['평가 단계','무엇을 측정할 것인가'],[
('분석','필요한 서술어·구성·높임·시제를 정확히 식별했는가.'),
('검색','필요한 근거가 후보 풀과 최종 전달 후보에 들어왔는가.'),
('적용 조건','문맥상 부적절한 후보를 선택하지 않았는가.'),
('생성 적용','제공한 적절한 지식을 올바른 구간·활용형으로 사용했는가.'),
('의미 보존','높임·시제·부정·역할·화행을 유지했는가.'),
('자연스러움','목표 지역·세대·관계에 맞는 표현인가. 방언 화자의 독립 검토 필요.'),
('일반화','오류를 보고 설계한 개발 문장 외의 새 자료에서도 개선되는가.')],[4,13.2])
p('단계별 증거를 모은 뒤에도 올바르고 간결한 계획을 줬을 때 실패한다면 그 사례를 생성 모델·추론 설정의 한계로 분리할 수 있다. 지금 단계에서는 분석과 후보 적합성 문제가 섞여 있어 모델 한계만으로 귀결하기 어렵다.')

h(chapters[12])
h('13.1 약 3분 보고 대본',2)
body('''“이번 연구는 문헌의 경상도 방언 지식을 그래프로 구조화하고, 표준어 문장에 필요한 지식을 찾아 생성 모델에 전달하는 방식입니다. 현재7종474쪽을 바탕으로3,761개 지식 진술을 만들었고, 원문 근거와 조건을 추적할 수 있습니다. Qwen의 입력 분석, Kiwi 형태 분석, KG 조회, Qwen 생성까지 연결되어 있습니다.

최근에는 같은 KG와 Qwen을 유지하면서 조회 방식과 전달 형식을 따로 비교했습니다. 기존 조회에 의미 검색을 추가한 조건, 같은 후보를 구간별 변환 계획으로 전달한 조건을 조합해50개 입력에서200개 출력을 확보했습니다. 후보 집합은37개 입력에서 달라졌지만 전체 번역 품질의 일관된 개선으로 이어졌다고 보기는 어렵습니다.

전달량도 실제 로그로 확인했습니다. 입력당 후보를 받은 구간은 평균3.0~3.2개, 구간당 후보는 약2.75개였습니다. 하지만 약23%의 구간은6개 상한에 도달했고 길이 제한 때문에 후보가 제외된 입력도 있습니다. 따라서 단순히 자료를 더 많이 넣거나 줄이는 것이 답이라고 결론 내리지는 않았습니다.

현재 확인된 문제는 단계별로 다릅니다. 파생 서술어가 검색 대상에서 빠지는 경우, 원문 예문을 KG에 놓친 경우, 보존 정보가 KG에 있지만 조회되지 않는 경우, 어미 목록을 실제 결합 규칙처럼 사용하는 경우가 있습니다. 강한 모델 출력을 참고하는 변환 조언자, 반례 검토자, 문헌 대조자와 조정자를 나누어 이를 검토했습니다.

검토에서 모델 제안9개와 새 대조 입력20개를 만들고 원문 인용22개를 확인했습니다. 동시에 앞에와 대라는 표현에 대한 기존 오류 판정도 정정했습니다. 다음 단계는 분석·조회 연결과 적용 조건을 보완하고, 문헌에서 확인한 누락을 추가한 뒤, 의미·높임 보존 검사를 통해 수정 효과를 분리해서 보는 것입니다. 아직 새 보완안을 기본 실행 코드에 반영하거나 품질 향상을 입증한 단계는 아닙니다.”''')
h('13.2 예상 질문과 답변',2)
table(['질문','보고할 답변'],[
('KG가 오히려 성능을 떨어뜨리는 것 아닌가?','일부 사례의 악화는 있지만 모델 간 비교만으로 KG의 순수 효과를 확정할 수 없다. 동일Qwen·동일KG의 통제 실험에서도 개선·회귀가 혼재한다.'),
('문헌 전체를 검토했다면서 왜 예문이 빠졌나?','페이지 검토 완료는 의미 완전 추출의 보증이 아니다. 최신 대조에서 특정 예문 누락을 확인했고 이를 보완 대상으로 기록했다.'),
('벡터 검색을 했는데 왜 문제를 못 잡았나?','검색 대상이 만들어지지 않는 분석 누락과 기존 필터를 공유하기 때문이다. 유사 후보 발견과 적용 조건 충족도 다른 문제다.'),
('노드 분기가 잘못된 것 아닌가?','표현·조건·근거를 분리하는 현재 구조는 유지 가능하다. 다만 목록/적용 규칙과 보존 지식의 사용 경계는 개선해야 한다. 전체 재구축 필요성은 아직 증명되지 않았다.'),
('모델에게 부족한 지식을 물으면 안 되나?','가능하며 제안9건을 만들었다. 원문 사실과 모델 가설을 구분하고 반례 검토 후 별도 실험해야 한다.'),
('왜 지금 기본 KG와 생성기에 적용하지 않았나?','최신 작업은 역할별 검토·통합과 실제 전달량 측정이었다. 제안과 반례 검토 결과를 만들었고 기본 반영은 다음 구현 단계다.'),
('24,000자씩 구간마다 넣는가?','아니다. 입력 전체 참고 패킷 예산이며 구간당 후보 상한6개와 별개다. 분석과 지시문을 포함한 전체 프롬프트는 더 길 수 있다.'),
('Neo4j가 매번 질의를 실행하는가?','현재 번역 조회는 로컬 facts/graph JSON을 Python으로 처리한다. Neo4j는 구조 확인용 별도 적재·시각화 경로다.'),
('현재 성공을 무엇으로 보고할 수 있나?','근거 추적 가능한 지식, 실행 가능한 연결, 통제된200개 출력, 실제 전달량, 구체적 누락과 잘못된 평가의 발견이다.'),
('다음 보고에서 무엇이 달라져야 하나?','필요 근거의 회수와 적용 조건, 생성 보존 오류가 수정 전후 어떻게 변했는지를 동일 기준으로 보여줘야 한다.')],[5.3,11.9])

h(chapters[13])
h('A.1 생성 system 프롬프트 원문',2)
code(protocol['config']['system_prompt'])
h('A.2 A/B 생성 user 지시문 원문',2)
job_a=next(j for j in jobs if j['condition']=='A')
code(job_a['request']['messages'][1]['content'].split('\n',1)[0])
h('A.3 C/D에 추가되는 계획 정책 원문',2)
job_c=next(j for j in jobs if j['condition']=='C')
body_c=json.loads(job_c['request']['messages'][1]['content'].split('\n',1)[1])
code(body_c['plan_policy'])
h('A.4 분석 프롬프트의 실제 지시 부분',2)
src=ROOT/'current_kg/experiments/run_auto_analysis_development.py';used.add(str(src.relative_to(ROOT)))
tree=ast.parse(src.read_text());fn=next(x for x in tree.body if isinstance(x,ast.FunctionDef) and x.name=='analysis_system')
prompt_literal=next(n.value for n in ast.walk(fn) if isinstance(n,ast.Constant) and isinstance(n.value,str) and n.value.startswith('표준어 발화의 입력 분석기다.'))
code(prompt_literal)
p('실제 함수는 이 뒤에 입력 특징 정의와 출력 JSON schema를 붙인다. 이 문서에서는 지시 부분을 그대로 싣고, 필드 의미는5장에 설명했다. 원문 전체 계약은 함수와 저장된 analysis_requests에서 확인한다. 기존38개는 캐시를 재사용했으므로 위 최신 프롬프트로 모두 새 분석했다는 뜻이 아니다.')
table(['분석 호출 설정','값'],[('모델','Qwen3.8-27B'),('temperature / top_p','0 / 1'),('max_tokens','4200'),('seed','20260923'),('thinking','false'),('response_format','json_object')],[5,12.2])
h('A.5 실제 DEV004 계획의 첫 후보',2)
p('아래는 저장된 요청에서 첫 후보만 발췌했다. 설명용 가상 예제가 아니다. 원래 그룹에는 후보3개가 있으며 다른 두 개는 -이라요와 -이래요다.')
candidate=plan['candidates'][0]
code({k:candidate[k] for k in ['id','source_pattern','target_pattern','conditions','claim_status','binding','evidence_pages','status']})
h('A.6 KG 작성과 조언자 역할의 입력·출력',2)
p('KG 작성은 문헌·작성 계약·담당 범위를 읽고 authored JSON과 페이지 검토 대장을 쓰는 에이전트 작업이었다. 단일 API 프롬프트 한 번으로474쪽을 일괄 변환한 것으로 설명하지 않는다. AUTHORING.md가 공통 출력 계약이며, 실제 문헌 사실과 분석자 해석을 구분하도록 요구한다.')
table(['역할','입력 구조','출력 구조'],[
('문헌 작성자','출처ID·담당 페이지·원문·표 이미지·작성 계약','source_id + facts; source_pattern/target_pattern/functions/conditions/constraints/scope/evidence/claim_status/usability/notes'),
('변환 조언자','실제 입력과 여러 모델 출력·현재 오류·KG 적용 한계','문제, 제안규칙, 적용조건, 보존사항, 제안예문, 반례, 불확실성, 출처 구분'),
('반례 검토자','제안 규칙·기존 실패 사례·문헌 대조 피드백','위험 일반화, 대조 입력, 보존 차이, 예상 검사, 실패 신호, 제안별 수용/보완'),
('문헌 대조자','제안 쟁점·KG·원문 텍스트/표·조회 기록','근거ID, 지원 상태, 페이지·정확 인용, 지원하는 것/지원하지 않는 것, 조치·한계'),
('조정자','세 결과와 원자료','검증 결과, 교차 정정, 통합 제안, 반영 상태, 해시, 통계')],[3.2,6.4,7.6])
p('위 역할 표는 실행 과제와 산출물의 요약이며 작업자에게 전달한 모든 메시지의 원문을 복원한 것은 아니다. 추적 가능한 실제 결과는 advisory 폴더의 각JSON에 있다.')

h(chapters[14])
p('다음은 최신 실험 outputs.json의200개 출력을 그대로 수록한 것이다. 문법·표기·띄어쓰기 오류를 교정하지 않았다. A/B/C/D의 차이가 없는 경우에도 조건별 원문을 유지했다. 전체 출력의 정답성이나 자연스러움을 보증하지 않는다.')
for row in analyses:
    sid=row['id'];h(sid,2);p('원문: '+row['input'])
    table(['조건','출력'],[(c,out_by[sid,c]) for c in 'ABCD'],[1.2,16])

h(chapters[15])
h('C.1 새 대조 입력: 아직 실행하지 않음',2)
p('기존50개와 문자열 중복이 없는20개 입력이다. 모델이 작성한 진단 자료이며 독립 평가셋이나 사람 정답이 아니다. 각 쌍에서 보존해야 할 차이를 먼저 명시했다.')
for issue in critic['items']:
    for pair in issue['minimal_pairs']:
        h(pair['pair_id'],2)
        table(['ID','표준어 입력'],[(r['id'],r['input']) for r in pair['inputs']],[3.5,13.7])
        p('보존할 차이: '+pair['must_preserve_difference'])
        p('위험한 일반화: '+issue['risky_generalization'])
h('C.2 조언자 제안의 상세 조건과 반례',2)
for a,name in zip(advisor['proposals'],adv_titles):
    h(a['id']+' '+name,2)
    p('문제: '+a['problem']);p('제안: '+a['proposed_rule'])
    p('적용 조건: '+' / '.join(a['applicability']))
    p('보존 사항: '+' / '.join(a['preservation']))
    p('반례·범위 제한: '+' / '.join(a['counterexamples_or_limits']))
    p('불확실성: '+a['uncertainty'])
    p('검토 판단: '+review_map[a['id']]['reason'])
    caption('출처 상태: model_advice_not_documentary_fact. 현재 생산 KG 적용 비활성. 제안 예문을 문헌의 확인형으로 취급하지 않는다.')

h(chapters[16])
h('D.1 이전 버전은 이 정도로만 구분',2)
table(['명칭','현재 보고에서 필요한 차이'],[
('legacy_graph','current_kg의 기존 KG와 기존 조회 방식. v2 비교는 KG만 고정적으로 교체한 실험이 아니라 전체 조회 시스템 비교.'),
('v2_direct','새KG의 직접 후보를 사용하되 활용 관계 확장 후보를 제외한 조건.'),
('v2_graph','활용 진술과 관계 경로를 추가한 조건. 단순히 같은 정보에 엣지만 추가한 순수 토폴로지 비교가 아님.'),
('최신A/B/C/D','종결 대응 보완 후의 같은KG v2에서 조회 방식×전달 형식을 비교. A가 해당시점 v2_graph 기준선.'),
('최신4역할 검토','출력·원문·KG를 대조한 진단과 제안 생성. 새로운 번역 실험 또는 기본 코드 변경이 아님.')],[4,13.2])
h('D.2 주요 근거 파일 지도',2)
evidence_groups=[
('현재KG 계약·규모','kg_v2/AUTHORING.md; fact.schema.json; source_registry.json; reports/build_summary.json'),
('현재 데이터·기본 코드','kg_v2/facts.jsonl; graph.json; sentence_retrieval.py; query.py; run_translation.py'),
('입력 분석','current_kg/experiments/run_auto_analysis_development.py; analysis_input.py; source_analysis_guard.py'),
('실험설계·고정 설정',EXP+'protocol.json; experiment.py; semantic_retrieval.py'),
('실제입력·분석·조회·출력',EXP+'jobs.json; analysis.json; retrieval.json; outputs.json'),
('통계',ADV+'measure_retrieval.py; retrieval_statistics.json; segment_lengths.tsv; input_lengths.tsv'),
('네역할 검토',ADV+'advisor.json; counterexamples.json; literature.json; reviewed_advice.json; coordinator_validation.json'),
('이전 평가 정정',ADV+'interpretation_corrections.txt; coordinator_report.txt'),
('모델비교','reports/astra_high_no_kg_20261012/protocol.json; reports/model_comparison_20261012/comparison.json'),
('이번 문서',str(OUT.relative_to(ROOT))+'/build_report.py; evidence_inventory.json; document_validation.json')]
table(['목적','저장소 경로'],evidence_groups,[4,13.2])
h('D.3 재현과 검증의 범위',2)
code('.venv/retrieval/bin/python reports/kg_v2_advisory_20261012/measure_retrieval.py\n.venv/retrieval/bin/python reports/kg_v2_advisory_20261012/integrate.py\n.venv/retrieval/bin/python reports/current_briefing_20261012/build_report.py')
p('측정·통합 스크립트는 기존 로그를 읽으며 Qwen을 새로 호출하지 않는다. 문서 작성 스크립트는 저장된 자료와 현재 코드의 지시문을 읽어DOCX와 도식을 만든다. 생산 실행 코드나 이전 실험을 덮어쓰지 않는다. 원본PDF·모델 캐시·접속 환경이 필요한 생성 재실행과는 구분한다.')
p('이번 문서 생성에서 실제 그래프 개수, 통계의 분모, 출력200개, 인용22개 검토 기록과 기존 파일 해시를 확인한다. 근거 파일의SHA256 목록을 evidence_inventory.json에 남긴다. API 주소·인증 정보는 문서에 넣지 않는다. Git에 변경사항이 있을 수 있으므로 커밋 하나만으로 최신 작업 상태를 대표하지 않는다.')
p('문서 검증 범위는 document_validation.json에 기록한다. Word/LibreOffice의 실제 페이지 렌더링이 불가능한 환경이면 그 한계를 명시한다. 내용 검증·DOCX 구조 검증과 실제 인쇄 페이지 검증을 동일하게 부르지 않는다.')

# Verify the report against frozen artifacts, independently of prose.
assert len(outputs)==200 and len(out_by)==200 and len(analyses)==50
assert len(graph['nodes'])==build['nodes']==25783
assert len(graph['edges'])==build['edges']==35358
assert len(audit['exact_source_quote_checks'])==22
for c in 'ABCD':
    rs=[r for r in seg_rows if r['condition']==c]
    assert len(rs)==s50[c]['delivered_segment_count_total']
    assert sum(r['candidate_count'] for r in rs)==s50[c]['candidate_count_total']
for pair in ['A_B','A_C','B_D','C_D']:
    a,b=pair.split('_')
    assert sum(out_by[sid,a]!=out_by[sid,b] for sid in input_by)==summary['output_difference_counts'][pair]
for path,digest in protocol['dependencies'].items():
    assert hashlib.sha256((ROOT/path).read_bytes()).hexdigest()==digest,path
    used.add(path)
for path in ['kg_v2/AUTHORING.md','kg_v2/fact.schema.json',ADV+'interpretation_corrections.txt',ADV+'measure_retrieval.py',EXP+'semantic_retrieval.py',EXP+'experiment.py']:
    used.add(path)
inventory={path:hashlib.sha256((ROOT/path).read_bytes()).hexdigest() for path in sorted(used)}
(OUT/'evidence_inventory.json').write_text(json.dumps({'as_of':'2026-10-12','files':inventory},ensure_ascii=False,indent=2)+'\n')
name='경상도_방언_KG_현재상황_상세보고_20261013.docx'
doc.save(OUT/name)
check=Document(OUT/name)
text='\n'.join(p.text for p in check.paragraphs)+'\n'+'\n'.join(c.text for t in check.tables for r in t.rows for c in r.cells)
assert all(x in text for x in ['3,761','25,783','35,358','4,133.71','2,742.57','22개','DEV030','LONG008','NEW012'])
validation={'status':'passed_structural_and_content_checks','file':name,
 'paragraphs':len(check.paragraphs),'tables':len(check.tables),'figures':len(check.inline_shapes),
 'text_characters':len(text),'source_files_hashed':len(inventory),'outputs_in_appendix':200,
 'new_model_calls':0,'production_files_unchanged':True,
 'actual_Word_LibreOffice_page_render_test':False,
 'limitations':['No Word/LibreOffice renderer is installed; page layout has not been independently rendered.',
 'Research findings are development observations, not independently certified dialect quality.'],
 'sha256':hashlib.sha256((OUT/name).read_bytes()).hexdigest()}
(OUT/'document_validation.json').write_text(json.dumps(validation,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(validation,ensure_ascii=False,indent=2))

if __name__=='__main__':pass
