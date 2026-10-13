"""Build the Korean research briefing from prose plus frozen project evidence.

Requires python-docx and Pillow. Does not call models or read credentials.
"""
import ast
from collections import Counter, defaultdict
import hashlib
import json
from pathlib import Path
import re
import sys

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
REPORT = ROOT/'kg_v2/reports/translation_20261012'
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.27), Inches(11.69)
section.top_margin = section.bottom_margin = Inches(.72)
section.left_margin = section.right_margin = Inches(.78)
section.header_distance = section.footer_distance = Inches(.3)
for name in ['Normal', 'Body Text', 'Caption', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    st = doc.styles[name]
    st.font.name = 'Noto Sans CJK KR'
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK KR')
doc.styles['Normal'].font.size = Pt(10)
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing = 1.18
for name,size in [('Title',27),('Heading 1',18),('Heading 2',13),('Heading 3',11)]:
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].font.color.rgb = RGBColor.from_string('173B55')
    doc.styles[name].paragraph_format.keep_with_next = True
doc.styles['Heading 1'].paragraph_format.page_break_before = True
doc.styles['Caption'].font.size = Pt(8)
doc.styles['Caption'].font.color.rgb = RGBColor.from_string('526572')
header = section.header.paragraphs[0]
header.text = '경상도 방언 KG 연구 | 연구 과정·구현·실험 브리핑'
header.style = doc.styles['Caption']
footer = section.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('2026-10-12 기준 · ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
settings = doc.settings.element
update = OxmlElement('w:updateFields'); update.set(qn('w:val'), 'true'); settings.append(update)
doc.core_properties.title = '경상도 방언 지식그래프 연구 상세 브리핑'
doc.core_properties.subject = '연구 이력, 모델 역할과 프롬프트, KG 구축 및 조회, 번역 실험, 높임 오류 진단'
doc.core_properties.author = '프로젝트 연구 기록 정리'

def read(path):
    return json.loads((ROOT/path).read_text())

def p(text, style=None):
    return doc.add_paragraph(str(text), style)

def h(title, level=1):
    par=doc.add_heading(title, level)
    name='sec_'+hashlib.md5(title.encode()).hexdigest()[:16]
    ident=str(len(doc.paragraphs))
    start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),ident);start.set(qn('w:name'),name)
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),ident)
    par._p.insert(0,start);par._p.append(end)

def body(text):
    for block in text.strip().split('\n\n'):
        p(block.strip())

def table(headers, rows):
    t=doc.add_table(rows=1, cols=len(headers)); t.style='Light Shading Accent 1'
    for cell,value in zip(t.rows[0].cells,headers): cell.text=str(value)
    repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for cell,value in zip(cells,row):
            cell.text=str(value)
            for par in cell.paragraphs:
                par.paragraph_format.space_after=Pt(3)
                for run in par.runs: run.font.size=Pt(9)
    p('')
    return t

def code(value):
    text=value if isinstance(value,str) else json.dumps(value,ensure_ascii=False,indent=2)
    for line in text.splitlines():
        par=p(line or ' ')
        par.paragraph_format.space_after=Pt(0)
        par.paragraph_format.line_spacing=1.0
        for run in par.runs: run.font.size=Pt(8)
    p('')

def figure(path, caption):
    path=ROOT/path
    with Image.open(path) as im: w,hh=im.size
    width=min(6.65,7.7*w/hh)
    par=doc.add_paragraph();par.alignment=1
    par.add_run().add_picture(str(path),width=Inches(width))
    p(caption,'Caption')

def literal_assignment(path,function,name):
    tree=ast.parse((ROOT/path).read_text())
    fn=next(n for n in tree.body if isinstance(n,(ast.FunctionDef,ast.AsyncFunctionDef)) and n.name==function)
    for n in ast.walk(fn):
        if isinstance(n,ast.Assign) and any(isinstance(t,ast.Name) and t.id==name for t in n.targets):
            return ast.literal_eval(n.value)
    raise ValueError((path,function,name))

def source_excerpt(path,function):
    text=(ROOT/path).read_text();tree=ast.parse(text)
    fn=next(n for n in tree.body if isinstance(n,ast.FunctionDef) and n.name==function)
    return ast.get_source_segment(text,fn)

p('경상도 방언\n지식그래프 연구','Title')
p('지금까지 무엇을 했고, 무엇을 확인했으며, 무엇이 남았는가','Subtitle')
p('브리핑 준비용 상세 설명서 · 2026년 10월 12일')
p('대상: 연구자 본인의 전체 과정 이해 및 다음 날 브리핑 준비')
p('기준 실험: KG v2 연결 후 38개 입력 × 4조건, 152개 최종 출력')
p('기준 저장소 커밋: 8eee9ce927c8e588081d7c39d769ef7fe52ef37b')
body('''이 문서는 실행 코드, 저장된 요청·응답, KG 원본, 검토 장부, Git 이력을 대조하여 작성했다. “설계했다”, “코드가 있다”, “실행했다”, “정확성을 확인했다”, “번역 품질이 향상됐다”를 각각 구분한다. 이전 설명의 혼동과 실패도 연구 과정에 포함한다.

현재 결과는 개발 실험이다. 문헌 전체를 검토했다는 말은 문헌의 모든 지식을 완벽하게 추출했다는 뜻이 아니며, 코드 검사를 통과했다는 말도 방언 화자가 자연스러움을 검증했다는 뜻이 아니다.

본문은 높임 후보 우선순위 수정 전의 38문장 결과를 기준으로 설명한다. 문서 작성 이후 수행한 수정이 있으면 마지막 부록에 별도 기록한다. 이전 실험 출력과 수치를 새 구현의 결과로 바꿔 적지 않는다.''')
p('읽는 순서: 1~3장으로 발표 흐름을 잡고, 4~15장으로 구현을 이해한 뒤, 부록의 실제 프롬프트·JSON·38문장 결과로 확인한다.')

h('목차와 증거 읽는 방법')
chapters=['1. 발표에서 먼저 말할 핵심','2. 연구 목적과 현재 범위','3. 시행착오와 버전 연표','4. 전체 처리 과정과 호출 경계','5. 모델·도구별 실제 역할','6. 원문 자료와 지식 구축','7. Astra 병렬 검토와 통합 검증','8. 기존 KG와 KG v2의 구조','9. 실제 입력 분석','10. 기존 조회 방식의 변천','11. KG v2 조회와 생성','12. 실험 설계·실행·비용','13. 결과와 오류 원인','14. 검증·평가의 경계','15. 재현과 파일 지도','16. 선행연구와 논문 구성','17. 다음 단계와 브리핑 질의응답','부록 A. 모델 프롬프트와 입출력 계약','부록 B. KG·분석·조회 실물 예시','부록 C. 38개 입력의 네 조건 출력 전체','부록 D. 근거 파일과 무결성','부록 E. 문서 작성 이후 높임 조회 수정']
for title in chapters:
    par=p('')
    link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),'sec_'+hashlib.md5(title.encode()).hexdigest()[:16])
    run=OxmlElement('w:r');props=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'175B83');props.append(color);run.append(props)
    text=OxmlElement('w:t');text.text=title;run.append(text);link.append(run);par._p.append(link)
body('''근거 표시는 [E01] 같은 번호를 사용한다. 부록 D에서 저장소 상대 경로와 파일 해시를 확인할 수 있다. “이력 문서의 기록”은 당시 보고를 뜻하며, 이번 문서 작성 중 그 실험을 새로 실행했다는 뜻이 아니다. 코드로 확인한 동작과 모델 출력에서 관찰한 현상도 구분한다.

한국어 설명에서 ‘높임’은 특히 두 가지를 구분한다. 청자 높임은 말하는 상대에 대한 대우이고, 주체 높임은 문장 속 행위 주체를 높이는 표현이다. ‘할머니께서 쉬고 계신다’에는 주체 높임이 있지만 청자에게 존댓말을 쓴다고 자동 판정할 수는 없다.''')

h('1. 발표에서 먼저 말할 핵심')
body('''연구 질문은 “문헌의 경상도 방언 지식을 구조화하고 입력에 필요한 지식을 제공하면, 로컬 LLM이 의미와 높임을 유지하면서 더 자연스러운 방언을 생성할 수 있는가?”이다. KG는 이 질문을 검증하기 위한 방법이며, 우월성이 이미 입증된 전제는 아니다. [E01, E03]

초기에는 악취 민원 대화를 여러 권역의 방언으로 합성하는 RAG 시스템을 만들었다. 이후 일반 표준어→경상도 방언 변환으로 초점을 옮겼다. 종결어미 중심 규칙 주입, 문헌 근거 중심 KG, 입력 분석·조건 조회, 전체 문헌 기반 KG v2 순서로 구조를 바꾸었다. 각 단계에서 지표의 순환성, 적용 위치 오류, 검색 누락과 높임 손실을 발견했다.

현재 KG v2는 7종·474쪽을 대상으로 작성·검토한 3,757개 진술을 담고 있다. 25,750개 노드와 35,308개 관계로 빌드했고, 기존 KG를 보존한 상태로 별도 Neo4j 공간에 적재·검증했다. Astra 작업자 3명이 남은 추출·검토를 분담했고 조정자가 통합과 원문 표본 검증을 수행했다. [E24~E29, E34]

입력 분석과 번역에는 Qwen3.8-27B를 사용한다. 그러나 실제 모델 사용은 이 두 문장으로 끝나지 않는다. Qwen의 문헌 대응 제안·검색 관련성 검토, 기존 조회기의 E5 임베딩, 과거 Claude 교란·평가 실험까지 역할을 분리해서 설명해야 한다. Astra의 문헌 작성은 연구 개발 시점의 작업이며, 문장 하나를 번역할 때마다 Astra가 호출되는 구조는 아니다.

현재 38개 입력의 네 조건 비교에서는 ‘먹다→묵다’와 일부 의문형이 개선됐다. 반면 높임 소실, 관형절 손상, 잘못된 활용형 조합이 남았다. 표준어 완전 유지가 기존 KG 8개에서 v2 5개로 줄었지만, 이 수치는 정확도나 자연스러움 점수가 아니다. 현재 성과는 추적 가능한 지식·조회 경로와 실제 연결·진단 결과이며, 전반적인 번역 품질 향상은 아직 입증하지 못했다. [E30~E33]''')
h('1.1 3분 발표 대본',2)
body('''“이 연구는 문헌에 있는 경상도 방언 문법과 어휘를 LLM에 어떻게 전달해야 의미와 말투를 보존한 방언 생성으로 이어지는지 살펴봅니다. 처음에는 검색한 어휘와 규칙을 프롬프트에 넣는 방식으로 시작했지만, 방언 표현이 늘어난 것과 번역 품질이 좋아진 것이 같지 않다는 문제가 드러났습니다.

그래서 원문 근거와 적용 조건을 갖춘 KG를 구축했습니다. 기존 KG 실험에서는 문장마다 필요한 어휘·활용 지식에 도달하는 경로가 부족했습니다. 이를 보완하려고 기존 자료를 보존하고 전체 7종 474쪽을 대상으로 KG v2를 새로 만들었습니다. 표준어 기본형, 방언 기본형, 실제 활용형과 조건을 분리하고, 모든 지식에 원문 근거를 연결했습니다.

최근에는 세 Astra 작업자가 남은 문헌 검토를 나눠 수행하고 조정자가 결과를 통합했습니다. 이후 Qwen 분석기와 생성기에 연결해 38개 입력을 네 조건으로 비교했습니다. 어휘와 일부 의문형은 개선됐지만, 높임과 문장 구조가 무너지는 사례도 나왔습니다. 현재는 KG의 양뿐 아니라 필요한 지식의 선택과 적용 통제가 중요하다는 것을 확인한 단계입니다. 특히 높임 어미가 KG에는 있으나 조회 정렬과 개수 제한 때문에 전달되지 않는 결함을 확인했습니다.

다음 단계는 그 결함을 고친 뒤 같은 개발 입력으로 회귀를 확인하고, 이후 고정된 설정과 별도의 문장·방언 화자 평가로 효과를 검증하는 것입니다.”''')

h('2. 연구 목적과 현재 범위')
table(['구분','현재 의미'],[
('최종 연구 대상','일반 표준어 발화→경상도 방언 발화. 단문과 여러 문장으로 된 발화를 함께 진단한다.'),
('방법 가설','문헌 지식을 형태·기능·조건·근거·관계로 조직하고 입력에 맞게 제공하면 생성에 도움이 되는가.'),
('학습 여부','현재 경로는 모델을 추가 학습하지 않는 분석·검색·프롬프트 주입 방식. KG 구축은 모델 가중치 학습이 아니다.'),
('지역·시대 정보','자료의 적용 범위를 보존한다. 세부 지역 분류나 지역별 생성 자체가 중심 목표는 아니다.'),
('악취 민원','초기 시스템의 도메인. 현재 38개 일반 발화 실험을 실제 악취 민원 데이터 평가로 부르지 않는다.'),
('AI-Hub','독립 평가 자료 후보. 문법 진술의 근거나 새 KG의 학습 정답으로 자동 편입하지 않는다.'),
('성과 기준','의미·부정·시제·높임·발화 의도 보존과 문법성·자연스러움. 단순 방언 표지 증가만으로 판단하지 않는다.')])
body('''KG와 RAG는 서로 배타적인 방법이 아니다. RAG는 외부 자료를 검색해 생성에 제공하는 흐름이고, KG는 그 외부 지식을 표현·조직하는 방법이다. 이 프로젝트에서는 KG 지식과 경로를 JSON 참고 자료로 바꿔 Qwen 프롬프트에 제공한다. 그래프 신경망을 학습하거나 Microsoft GraphRAG의 커뮤니티 요약을 실행한 것은 아니다.

현재 생성기는 목표를 넓게 ‘경상도 방언’으로 받는다. 특정 화자의 시대·지역·관계·격식을 명시적으로 고정하지 않으면, 원문에 맞는 하나의 말투를 선택하는 문제가 남는다. 문헌의 지역·시대 범위를 기록한 것만으로 실제 생성에서 그 범위가 자동 보장되지는 않는다.''')

h('3. 시행착오와 버전 연표')
table(['시기·명칭','실제로 한 일','발견과 다음 선택'],[
('초기~2026-08 / archive/v1_retrieval','악취 시나리오 생성, S1 직접 방언 생성과 S2 표준어 생성 후 방언 변환. 무지식·예시·어휘/규칙 검색 등 비교.','어미 표지만으로는 어휘 활용이나 의미 보존을 설명할 수 없었다. 사후 유리한 지표 선택과 순환 평가를 경계하게 됐다.'),
('2026-08 / archive/v2_failed','종결어미 39규칙·음운 14항목, 자질 추출과 규칙 매칭, plain/KG/전체 어미 덤프 실험.','그래프라는 명칭과 달리 주요 조회는 사전 필터였다. 치환 오류, 어미 강제, 입력/평가 누수와 지표 문제가 드러났다.'),
('2026-08-31 감사','이력 문서에 측정 실패를 확정. copy와 상수 기준선을 재검토.','기존 수치로 KG의 양·음 효과를 주장하지 않기로 했다. 평가와 지식 근거를 다시 설계했다.'),
('2026-09 문헌·Claim 검토','원문 인용, 출처 정책, 원자 주장 초안, 공통 개념과 적용 조건을 검토.','문헌 설명·저자 해석·우리의 해석을 분리. 이전 Claim 프로토타입은 현행 구조와 합산하지 않음.'),
('2026-09~10 / current_kg','문헌 근거 기반 기록·Assertion·표현·조건 그래프 구축. 최종 334기록·6,214노드·15,978관계.','의문형·높임·인용·연결 구성의 적용 위치와 조건을 보강했지만 어휘→활용 경로가 제한적이었다.'),
('2026-09-23 개발 입력 시작','연구자 구성 표준어 30개 무KG 출력, 이후 시드 고정 무KG/수동 KG 30쌍=60출력.','Codex가 이전 출력도 본 상태에서 시작 관계를 수동 선택했다. 자동 조회 시스템·독립 평가·그래프만의 효과로 해석하지 않음.'),
('2026-10-07~08','Qwen이 문헌 1,064단위의 1차 대응 제안을 수집. 원문과 대조한 일부 제안을 채택·수정.','제안 수집 완료와 의미 승인 완료를 구분. 잘못된 ID와 인용을 격리하고 원응답을 보존.'),
('2026-10-08 조회 실험','문자·E5·혼합 검색, 관계 확장, 30문장×4조건 비교. 미통합 경로를 발견한 뒤 조회기를 통합하고 별도 재실행.','이전 120개와 통합 후 120개를 구별. 통합 C/D는 모든 30개 출력이 같았음.'),
('2026-10-08 긴 발화','8개 긴 발화×4조건=32개 추가 출력.','긴 입력만으로 문제가 해결되지 않음. 먹고→묵고 경로 부재와 의문 분석 오류 등을 확인.'),
('2026-10-12 / kg_v2','7종 전체 원문에서 새 구조로 2,379진술 초안 작성. 세 Astra 작업자가 남은 항목 보완.','1,378개 추가, 기존 759개 수정·보강. 최종 3,757진술. 전체 독해와 완전 추출을 구분.'),
('2026-10-12 연결·진단','Qwen 분석→Kiwi→v2 조회→Qwen 생성 연결. 38개×4조건 비교 및 Neo4j PNG.','어휘 개선과 높임·관형절 회귀 혼재. -예 후보의 정렬·선택 누락을 후속 진단으로 확인.')])
body('''용어 주의: archive/v1_retrieval은 초기 민원 RAG이고, archive/v2_failed는 8월의 실패 실험이다. current_kg는 그 뒤 다시 구축한 기존 문헌 KG이며, 10월 12일 비교에서 legacy_graph로 부른다. 현재 kg_v2는 별도의 새 지식 표현이다. “v1 대 v2”라고만 말하면 이 네 작업을 혼동하게 되므로 발표에서는 경로와 시점을 함께 말한다. [E02, E06, E07]

8월 실패에서 얻은 교훈은 지금도 유효하다. 끝음절 한 글자 일치 점수에서는 입력 복사가 34/60, plain이 30/60, KG가 23/60, 덤프가 26/60이었다는 감사 기록이 있다. 이것은 당시 지표의 부적합성을 보여주는 기록이지, 현재 KG v2의 정확도가 23/60이라는 뜻이 아니다. 방언은 여러 정답 표현이 가능하고, 한 글자 지표는 나머지 문장의 의미·문법 손상을 보지 못했다. [E07]

초기 RAG 이력에는 few-shot이 어미축에서, 지식 주입이 어휘축에서 다른 양상을 보였다는 해석이 있다. 당시 수치와 해석은 해당 데이터·모델·지표에 한정한다. 본 문서는 그 실험을 새로 재채점하지 않았으며, 과거의 ‘확정’ 표현을 현재 시스템의 독립 품질 검증으로 승계하지 않는다. [E06]''')

h('4. 전체 처리 과정과 호출 경계')
figure('reports/briefing_20261012/process.png','그림 0. 지식 구축과 새 문장 번역의 경계. 38문장 개발 비교에서는 분석을 캐시에서 공유했고, Neo4j는 별도 확인·시각화 경로다.')
h('4.1 연구 개발 시점: 지식 구축',2)
code('원본 PDF 7종\n  → pdftotext로 페이지별 추출 + 해시\n  → Astra가 본문·표·각주 해석, 필요한 페이지 이미지 대조\n  → authored/*.json 진술 + reviews/*.json 페이지 대장\n  → Astra 작업자 3명의 남은 추출·검토 + 조정자 통합\n  → 스키마·인용·페이지·ID·조건·보존 검사\n  → facts.jsonl + graph.json\n  → 구조 진단 / 문장 조회 / Neo4j 적재·그림')
body('''PDF 추출과 그래프 빌드는 Python·명령행 도구가 수행한다. 문헌의 의미를 읽고 진술을 작성하는 것은 Astra의 작업이다. 작성 스크립트가 있다는 사실을 “코드가 자동으로 문헌 의미를 이해했다”라고 설명하면 안 된다. Astra가 여러 단계의 도구 호출·파일 작성을 하는 에이전트 작업이지, 474쪽 전체를 단일 HTTP 요청으로 보냈다가 완성 KG 하나를 받은 형태가 아니다. [E12, E14]''')
h('4.2 새로운 문장 한 개를 번역할 때',2)
code('표준어 입력\n  → Qwen 입력 분석 호출: JSON\n  → 코드가 원문 구간 결속·스키마 검사\n  → Kiwi 형태소 분석과 원문 기반 보정\n  → Python KG v2 조회: 조건부 진술·활용 관계\n  → Qwen 방언 생성 호출: 번역문 문자열\n  → 요청·응답·분석·조회 결과 저장')
body('''정상적인 새 입력 경로는 Qwen 분석 1회와 생성 1회다. 분석 검증에 실패하면 재시도가 추가될 수 있다. Kiwi와 KG 조회 자체는 원격 생성 모델 호출이 아니다. 같은 문장을 네 조건으로 실험할 때는 분석을 공유하고 생성만 조건별로 호출한다. 이번 38문장 비교는 저장된 분석을 재사용했으므로 38개의 분석을 모두 새로 요청한 실험이 아니다. 별도의 live smoke에서 새 분석부터 끝까지 실행한 증거를 보존했다. [E17~E19, E30]

현재 번역의 KG 조회는 로컬 facts.jsonl과 graph.json을 Python으로 읽는다. Qwen이 Cypher를 만들거나 매 문장마다 Neo4j 서버를 조회하는 구조는 아니다. Neo4j는 동일한 그래프를 시각화·확인하는 별도 경로이며, 적재된 노드·관계·속성이 원본과 일치하는지 검증했다.''')

h('5. 모델·도구별 실제 역할')
table(['모델/도구','입력 → 출력','사용 상태'],[
('Astra / gpt-6-astra','원문·작성 계약·기존 진술·검토 과제 → 출처별 진술 JSON, 페이지 대장, 수정·검토 보고','KG v2 작성·검토·통합에 실제 사용. 문장 번역 런타임의 필수 호출은 아님.'),
('Qwen3.8-27B / 입력 분석','표준어 발화 + 입력 특징 정의·JSON 스키마 → clauses/source_units JSON','새 입력 translate에서 실제 호출. 비교 실험에서는 저장 분석 공유.'),
('Qwen3.8-27B / 방언 생성','표준어 + 분석 + 형태소 + 조건별 참고 자료 → 방언 발화 문자열','현재 네 조건의 생성 모델.'),
('Qwen / 문헌 1차 대응','문헌 단위·출처·후보 KG 설명 → 유형·요약·원문 용례·후보 ID·한계 JSON','10월 7~8일 1,064단위 제안 수집. 현행 v2 자동 작성 승인과 다름.'),
('Qwen / 검색 관련성 검토','입력 구성·문자 검색 후보 → 선택 ID·입력 단위 ID·직접 인용·관련성 이유','기존 source_retrieval 확장 경로. 현재 v2 조회기는 이 호출을 사용하지 않음.'),
('multilingual-e5-small','query:/passage: 접두어를 붙인 텍스트 → 정규화 임베딩 벡터','기존 혼합 검색에 실제 사용. 현재 v2 직접 조회에는 미사용.'),
('Claude / 과거 실험·검토','과거 생성물·평가 지시 → 표준어화/타권역화 교란, 평가·설계 검토','archive 및 CLAUDE_REVIEW_LOG의 역사적 사용. 새 KG 독립 검증 완료로 간주하지 않음.'),
('Kiwi','표준어 문자열 → 형태소·품사·원문 위치','로컬 분석기. Astra/Qwen API 호출과 별개이며 분석 오차가 있음.'),
('Python·jsonschema·Graphviz·Neo4j','구조화 파일·그래프 → 검사 결과·조회·PNG','모델이 아닌 코드·DB·렌더링 도구.')])
body('''“모델 두 개를 쓴다”는 설명은 KG 작성자와 최종 번역기만 말하는 요약이다. 실제 실험을 이해하려면 모델 종류, 같은 모델의 역할, 역할별 호출 수, 캐시 재사용 여부를 따로 기록해야 한다. Qwen 분석·문헌 대응·번역은 같은 모델명이어도 목적과 프롬프트·출력 계약이 다르다.

Qwen3.8-27B는 이 프로젝트의 요청 및 서버 응답 모델명이다. 이 이름만으로 정확한 공개 체크포인트, 학습 이력, 현재 양자화·서빙 설정이 불변임을 증명할 수 없다. 과거 문서의 Qwen3.6 또는 FP8 표기를 현재 실험의 확인된 설정으로 덮어쓰지 않는다. Astra 역시 여기서는 작업 로그의 모델 ID를 기록한 것이며, 이 문서가 제품의 현재 공개 제공·요금·벤치마크를 검증한 것은 아니다.

임베딩은 intfloat/multilingual-e5-small, revision 614241f622f53c4eeff9890bdc4f31cfecc418b3를 로컬 CPU에서 사용했다. 문서 조각은 최대 384토큰, 이동 간격 320토큰으로 나누고, passage:/query: 접두어와 최대 512토큰 입력을 사용한다. 벡터와 문자 검색을 순위 기반으로 합친다. 이 모델이 방언 검색에 최적이라고 증명한 것은 아니다. [E23]''')

h('6. 원문 자료와 지식 구축')
summary=read('kg_v2/reports/build_summary.json')
registry=read('kg_v2/source_registry.json')
stats={s['source_id']:s for s in summary['sources']}
table(['자료 ID','문헌','PDF 쪽','진술'],[(s['source_id'],s['title'],stats[s['source_id']]['pages'],stats[s['source_id']]['facts']) for s in registry['sources']])
body('''대상은 학술 문헌 3종과 동남 방언 관련 문서 스냅샷 4종이다. 스냅샷에는 재인용·편집 설명이 있으므로 학술논문과 같은 증거 강도로 일괄 취급하지 않는다. 문헌 간 의존성을 남겨 같은 원전을 여러 번 독립 근거로 센 것처럼 보이지 않게 한다. 페이지는 PDF의 1부터 시작하는 페이지 번호이며 인쇄 쪽수와 다를 수 있다. [E42]

474쪽에는 표지·목차·서지·빈 페이지도 포함된다. 각 페이지의 검토 상태를 기록하되, 모든 페이지에서 생성 가능한 규칙이 나와야 한다고 강제하지 않는다. 애매한 표·전사·저자 가설·통계는 참고용 또는 한계로 보존한다. 650개 어휘 대응과 675개 활용 진술도 독립 어휘 수나 보편적인 생산 규칙 수가 아니다. 출처·조건·표 셀 단위가 포함된다.

추출은 PDF 텍스트를 페이지별로 보존하고 필요 시 이미지로 행·열·각주를 확인하는 방식이다. 묵고가 있는 표의 줄바꿈을 잘못 읽지 않도록 원본 이미지를 대조했다. 삐다/뿌다 활용표의 ‘삐가/삐가’ 중복 셀은 임의로 ‘뿌가’로 고치지 않고 불확실한 자료로 남겼다. 원문이 이상하다고 모델이 자연스럽게 고쳐 적으면 원문 근거 추적이 깨진다.

자료의 역할을 분리했다. 별도 어휘 JSON은 후속 후보이며 이번 7종에 포함하지 않았다. AI-Hub는 독립 평가 후보이고, 분석 연결을 설명하는 참고 문헌은 방언 지식의 직접 근거가 아니다. 원본 PDF·말뭉치 파생 문장쌍·인증정보는 Git 공개 대상에서 제외한다.''')
h('6.1 원문에서 그래프까지의 산출물',2)
table(['단계','파일','검사/의미'],[
('등록','source_registry.json','출처 ID, 제목, PDF 경로·해시, 허용 범위와 의존성'),
('추출','local/pages.jsonl','PDF별 페이지 텍스트. 로컬 보존'),
('모델 작성','authored/*.json','원문에 근거한 지식 진술의 수정 원본'),
('페이지 검토','reviews/*.json','페이지 분류, 진술 ID, 확인한 내용, 남은 gap'),
('빌드','facts.jsonl, graph.json','스키마·인용·ID 검사 후 생성한 통합 산출물'),
('감사','reports/integration_audit.json 등','기존 ID·원문 근거·페이지 처리·보존 여부'),
('문장 실험','reports/translation_20261012/*','분석·조회·요청·출력과 품질 관찰')])

h('7. Astra 병렬 검토와 통합 검증')
body('''첫 초안은 전체 문헌 범위를 읽고 작성한 2,379개 진술이었다. 그 상태를 단순히 ‘KG v2 완료’라고 표현하면 남은 추출·검토까지 끝났다는 오해가 생긴다. 이후 사용자가 Astra 작업자 3개를 명시적으로 요청했고, 기존에 공백이 기록된 162쪽을 분담해 보완했다. 이번 병렬 작업은 처음부터 문헌 전체를 세 모델이 각각 독립 구축한 삼중 평가가 아니다. [E03, E25]

로컬 세션 메타데이터에서 세 작업자의 model=gpt-6-astra, reasoning_effort=medium과 부모 세션을 확인했다. 실제 분담은 문헌 7개에 모델 7개를 붙인 것이 아니라, 아래 세 묶음이다. 부모 조정자가 공통 스키마와 통합 빌드·검증을 맡았다. 같은 모델의 여러 작업 문맥이 병렬로 진행된 것이며 모델 가중치를 합치거나 다수결로 문법 정답을 판정한 것이 아니다.''')
table(['작업자','범위','원래 공백 / 추가 진술','핵심 역할'],[
('jo_review','조현준 2018','66쪽 / 600개','연구사·지역·높임 분류표, 종결어미와 빈도·원문 모순 보존'),
('grammar_honorific_review','문법·상대높임법 스냅샷','59쪽 / 528개','활용표 셀, 어휘 각주, 표기형/발음형, 결합·범위 근거 보강'),
('other_sources_review','함양·개요·음운·-거로','37쪽 / 250개','높임 범주 과잉 해석 수정, 발화 대응과 음운 범위 보강'),
('root 조정자','공통 스키마·통합·교차 점검','추가 총 1,378개','ID·페이지 대장 통합, 출처 문자열 검사, 표본 이미지/의미 대조와 정정')])
body('''결과는 2,379→3,757개 진술이며 기존 759개가 수정·보강됐다. 원래 공백 162쪽은 해결 84쪽과 원문·자료 한계 78쪽으로 정리했다. 새로 발견한 문법 17쪽의 불확실한 표 문제를 더하면 현재 제약이 남은 페이지는 79쪽이다. ‘78’과 ‘79’는 분모와 시점이 달라 모순이 아니다.

통합 검증은 두 층이다. 기계 검사는 전체 진술의 스키마, ID, 인용 문자열의 페이지 내 존재, 그래프 끝점, 페이지 대장, 해시를 검사했다. 의미 검토는 조정자의 표본 원문·이미지 대조와 작업자 간 교차 검토다. 예를 들어 묵다의 각주와 활용표를 함께 확인했고, 함양의 특정 -십시오 부재를 높임 범주 전체 부재로 넓힌 표현을 수정했다. 일부 발화 대응에서 잘못 연결된 인용도 바로잡았다. [E26]

공통 작업 계약은 AUTHORING.md에 남아 있다. 작업자에게 전달한 개별 지시문의 정확한 본문은 로컬 세션 로그에서 암호화되어 있어 이 문서에 원문이라고 재현하지 않는다. 부록 A에는 확인 가능한 작성 계약과 분담을 바탕으로 재구성한 작업 지시의 구조를 구분해 적었다. 같은 Astra의 작성·검토는 사람 전문가나 다른 모델의 독립 언어학 평가가 아니다.''')

h('8. 기존 KG와 KG v2의 구조')
table(['항목','기존 current_kg','새 kg_v2'],[
('지식 중심 단위','문헌 기록 Record 안의 표현·관계·Assertion','명시적인 KnowledgeStatement와 출처별 source/target 표현'),
('주요 관심','어미·기능·결합·적용 위치와 문헌 근거','어휘 대응·방언 기본형·실제 활용형까지 이어지는 전달 구조'),
('조건','입력·후보·목표의 특징과 적용 위치','owner=input/candidate/target/context를 명시'),
('원문 추적','기록·주장에 연결된 Evidence','진술→Evidence→SourcePage→Source'),
('관계 활용','관련 기록·제약·인용·활용 설명 등을 확장','어휘 대응 진술→활용 진술의 candidate_has_inflection'),
('규모','334기록 / 6,214노드 / 15,978관계','3,757진술 / 25,750노드 / 35,308관계')])
body('''334기록과 3,757진술은 집계 단위가 달라서 “지식이 정확히 11배 늘었다”라고 말하면 안 된다. 그래프가 커진 일부 이유는 조건·근거·표현을 별도 노드로 표현했기 때문이다. 노드 수 증가는 정확도나 검색 포괄률의 지표가 아니다.

‘카드’는 KG의 별도 노드 종류가 아니다. 생성 모델에 보낼 지식·조건·원문 구간·근거를 묶은 JSON 참조 항목을 편의상 카드라고 불렀다. 원래 노드·관계는 그래프 안에 있고, 카드는 그 일부를 모델이 읽도록 직렬화한 것이다. 따라서 “카드가 같다”는 것은 비교 조건이 같은 참고 항목을 보냈다는 뜻이다. 서로 다른 30개 입력이 모두 같은 노드를 조회했다는 뜻은 아니다.''')
counts=read('kg_v2/reports/neo4j/schema_counts.json')
meanings={'Source':'문헌','SourcePage':'PDF 페이지와 검토 상태','KnowledgeStatement':'하나의 출처 기반 지식 진술','ExpressionUse':'해당 진술의 출발/목표 표현','Constraint':'자유 서술 제약','Evidence':'원문 인용과 페이지','Function':'문장·문법 기능','Condition':'판별 주체·필드·연산·값'}
table(['노드 유형','수','설명'],[(k,v,meanings[k]) for k,v in counts['nodes'].items()])
figure('kg_v2/reports/neo4j/schema.png','그림 1. 실제 Neo4j 조회 결과를 유형별로 집계해 Graphviz로 그린 구조도. 개별 25,750개 노드를 한 화면에 모두 펼친 그림이나 Neo4j Browser 캡처는 아니다.')
h('8.1 먹다→묵다→묵고를 따라 읽기',2)
body('''첫 진술은 표준어 기본형 먹다와 방언 기본형 묵다의 대응을 담는다. 두 번째 진술은 방언 후보가 묵다이고 요청 어미가 -고일 때 문헌 표에서 묵고가 확인된다는 정보를 담는다. candidate_has_inflection은 이 두 진술을 연결한다. 표준어 입력이 처음부터 ‘묵다’여야 한다고 요구하는 잘못된 조건 배치를 피하려고, 활용 조건의 주체를 input이 아니라 candidate로 분리했다.

연결은 후보를 찾는 길이다. 연결된 두 진술의 지역·뜻·시대가 자동으로 모두 같아지는 것은 아니다. 현재 44개의 어휘 대응 진술에 484개의 활용 관계가 있다. 명사처럼 활용 대상이 아닌 어휘도 있으므로 44/650을 동사 활용 포괄률로 계산해서는 안 된다. 표에 없는 어미는 확인형으로 만들어 넣지 않는다.''')
figure('kg_v2/reports/neo4j/eating_path.png','그림 2. 실제 먹다→묵다→묵고 부분 그래프. 조건과 원문 근거가 함께 연결된다.')
h('8.2 조건·상태의 의미',2)
table(['필드','의미·실제 동작'],[
('source_pattern / target_pattern','출발·목표 표현, 기본형, 의미, 특징. domain은 standard/dialect/shared/metalinguistic.'),
('functions','declarative, interrogative, cause 등. 문헌 분류와 입력의 기능을 연결하는 단서.'),
('conditions','owner + field + op + value. 코드가 satisfied/failed/unknown을 판정한다.'),
('constraints','require/exclude/scope/uncertainty의 자유 서술. 모든 내용을 코드가 의미적으로 판정하지는 못한다.'),
('scope','지역·높임체·사용역·시기·추가 메모. 기록과 실제 통제는 다른 문제다.'),
('claim_status','실제 형태·예문, 원문 일반화, 저자 가설, 이차 보고, 작성자 해석, 부정 예문을 구분.'),
('usability','candidate / reference_only / blocked. candidate는 조회 후보이며 생성 품질 승인 아님.'),
('evidence','pdf_page와 실제 인용. 문자열 존재와 의미적 뒷받침의 충분성은 별도.')])

h('9. 실제 입력 분석')
body('''Qwen 분석 프롬프트는 번역하거나 방언 후보를 제안하지 말고 입력의 문법 정보를 JSON으로 출력하라고 요구한다. 목표 방언 정답이나 KG 후보 ID를 주지 않는다. 원문 의미를 보고 문장 형식과 발화 의도, 설명/판정 질문, 주체·청자 높임, 시제·선어말어미, 주요 구성과 검색용 설명을 작성한다. [E18]

clauses는 주절(main)과 연결절(connective)을 표현한다. source_units는 인용·관형절·보조용언·조사·관용 구성 등 더 자유로운 입력 구성을 보존한다. text는 원문의 연속 부분 문자열이어야 하고, 같은 문자열이 여러 번 나오면 occurrence로 구분한다. span과 ID는 코드가 계산한다. 모델이 문자 위치를 임의로 만들어 넣는 방식이 아니다.

스키마 검사와 원문 결속 검사는 허구의 구간이나 형식 오류를 막지만 문법 해석의 정답을 보장하지 않는다. 모델이 설명 질문을 판정 질문으로 잘못 분석하거나 중단과 동시 행동을 혼동할 수 있다. 자유 설명은 검색 가설이고 확정 조건으로 바로 승격하지 않는다.

Kiwi 원문 보정은 선어말과 높임의 관찰값을 다시 구성한다. 이 보정도 보수적인 지원 목록에 의존한다. 예를 들어 힘드셨죠?의 높임이 unknown으로 남는 데에는 Qwen만이 아니라 후속 원문 보정기의 종결형 지원 범위가 관여한다. 앞서 이 문제를 단순히 ‘Qwen 분석 오류’라고 표현했다면 불완전한 설명이다. 최종 분석은 모델 출력과 코드 보정의 합성 결과다. [E19]

v2의 추가 semantic_guard는 왜·어디·어떻게·언제·몇 같은 명시적 의문 단서로 일부 질문 유형을 수정한다. 선택 질문은 제외한다. 이번 LONG001의 왜 아직 안 와?를 수정했고, 공정한 비교를 위해 기존 KG 조건도 같은 수정 분석으로 다시 조회했다. 이 좁은 보정으로 모든 질문 의미가 해결된 것은 아니다.''')

h('10. 기존 조회 방식의 변천')
body('''기존 조회에는 여러 개발 경로가 있었으므로 하나로 묶어 설명하면 잘못된다. 초기에는 입력 특징으로 문법 기능과 조건을 맞췄다. 다음 확장에서는 원문·구성 설명을 문자 검색하고 Qwen이 관련 후보를 검토했다. 이후 E5 임베딩과 문자 검색을 같은 문서 조각에 적용해 혼합 검색을 비교했다.

10월 8일의 controlled_development 경로는 혼합 검색 상위 기록을 전달했지만 기존의 세밀한 기능/조건 조회를 모두 통합하지 않았다. 이 불일치를 발견하고 integrated_retrieval로 기능·원문 형태·혼합 검색과 원문 구간 결속·조건 평가를 통합했다. 따라서 수정 전 120개 결과와 통합 후 120개 결과를 별도 실험으로 보존했다. [E03, E09, E22, E36, E37]

통합 조회기는 입력의 각 구성별 후보를 찾고, 관련 관계를 최대 2단계 확장한다. has_applicability_constraint, has_inflection_description, has_construction_component, shares_quotative_construction, same_usage_candidate, contrasts_in_function 등 제한된 관계만 따른다. 도착한 후보도 자기 입력 구간의 조건을 다시 검사하며 부모의 특징을 임의로 상속하지 않는다.

잘못된 적용 위치, 조건 미일치·미확인, 검토 대기, 미해결 예외, 종결어미의 높임 조건 부재 등을 제외한다. 참고 자료는 구간당 최대 6개, 전체 24,000자 범위다. 이 숫자는 문법 정답 수가 아니라 프롬프트 구성의 제한이다.

이 구조에서도 30문장의 record_hybrid와 graph_hybrid가 같은 후보를 선택했고, 전체 요청은 29개에서 같았다. 출력은 30개 모두 같았다. 관계가 존재한다는 것과 해당 입력에서 새롭고 적절한 지식이 추가된다는 것은 다르다는 결과다. [E37]''')

h('11. KG v2 조회와 생성')
h('11.1 후보를 찾는 방법',2)
body('''형태소 분석에서 기본형과 실제 원문 위치를 얻는다. lexical_mapping/phonology 중 출발 표현이 standard 또는 shared인 진술을 기본형과 맞춘다. 예컨대 먹고의 먹다를 발견하면 먹다→묵다를 후보로 찾는다. 연결된 활용표를 조회할 때는 요청 어미도 맞춘다. [E15, E16]

현재 완성 활용 경로는 -고, -아/어, -아/어라, -(으)라고, -(으)니 등 구현된 어미 집합에 한정된다. 선어말이 있으면 단순 표 셀을 그대로 덮어쓰지 않도록 경로를 보류한다. 그래서 묵었어·묵자 같은 출력이 보였다고 모든 완성형이 KG 활용표에서 조회됐다고 주장할 수 없다. 기본형 대응을 보고 Qwen이 활용한 경우를 분리해야 한다.

어미·구문은 원문 어미와 source_pattern의 대응 또는 문장 기능으로 찾는다. 기능만으로 찾는 경로는 metalinguistic 설명형 종결어미로 제한했다. 실제 표준어 형태가 명시된 규칙을 그 형태가 없는데 기능만 같다는 이유로 전달하던 초기 오류를 수정한 것이다.

candidate 상태이며 구조화 조건이 모두 확인된 후보만 전달한다. blocked/reference_only/부정 예문과 unknown 조건은 제외한다. 그러나 자유 서술 제약·지역·시대·세부 의미는 완전히 자동 판정되지 않는다. 현재도 최종 선택의 일부를 Qwen에게 맡긴다.''')
table(['조건','생성에 전달되는 자료','검증하고 싶은 효과'],[
('no_kg','원문·공통 분석·형태소. references=[], graph_paths=[]','KG 지식이 없는 기준선. 분석도 없는 순수 원문 기준선은 아니다.'),
('legacy_graph','기존 통합 조회기의 진술·관계 패킷','기존 KG＋기존 조회 전체 시스템'),
('v2_direct','v2에서 선택한 진술. 완성 활용형 진술 제외','새 표현·후보 선택으로 전달되는 지식'),
('v2_graph','v2 진술＋연결된 활용 진술＋실제 관계','확인된 활용 지식을 추가 제공하는 효과')])
body('''v2_direct와 v2_graph는 같은 KG를 사용하는 두 실험 조건이다. 내부 후보 발견 과정은 공유하지만, 직접 조건은 완성 활용 진술을 보내지 않는다. 관계 조건은 경로의 양 끝 진술이 실제 패킷에 들어간 경우만 관계를 전달한다. 이것은 전체 그래프의 임의 다중 홉 검색을 켠/끈 비교가 아니다.

같은 입력 구간에 최대 6개 참조를 보내며, 원문 근거·범위·검사 결과를 포함한다. 기본형 대응 우선순위는 100, 활용 경로는 110, 원문 어미 일치는 85, 기능 경로는 60을 기본으로 삼고 조건 수를 일부 가산한다. 현재 기준 실험은 동점에서 ID를 사용했다. 이 마지막 선택이 높임 후보 다양성을 해치는 결함을 만들었다.

생성 프롬프트는 조건에 맞는 자료만 선택하고 의미·부정·시제·높임·인용·발화 의도를 유지하라고 지시한다. 하지만 이는 자연어 지시다. 허용된 위치와 활용형만 생성하도록 강제하는 디코더도, 생성 뒤 자동으로 높임을 검증·재생성하는 모듈도 현재 기준 실험에는 없다. [E17]''')

h('12. 실험 설계·실행·비용')
table(['실행','입력·출력','해석 경계'],[
('10/08 혼합 검색 개발','30개×4조건=120출력','기능·조건 경로 미통합 개발 진단'),
('10/08 통합 조회 개발','30개×4조건=120출력','C/D 출력 동일 30/30. 독립 평가 아님'),
('10/08 긴 발화','8개×4조건=32출력; 분석 시도 13회 기록','단문 한계를 보기 위한 추가 개발 입력'),
('10/12 v2 첫 실행','38개×4조건=152출력; 생성 시도 154회','조회 수정 전 결과 보존'),
('10/12 v2 최종 비교','38개×4조건=152출력','기준선 76개 재사용＋수정 v2 출력 76개. 신규 생성 시도 80회')])
protocol=read('kg_v2/reports/translation_20261012/protocol.json')
table(['설정','값'],[(k,v) for k,v in {'model':protocol['config']['requested_model'],'seed':protocol['seed'],**protocol['config']['generation'],'반복':'1개 시드·1회 반복','참고 자료 한도':'24,000자; 토큰 수가 아님','입력':'짧은 30개 + 긴 발화 8개','공유 분석':'기존 Qwen 분석 + Kiwi 형태 분석 + 제한적 공통 보정'}.items()])
body('''최종 보고서의 actual_attempts=158은 재사용 기준선의 시도 이력까지 합친 실행 계보다. 최종 단계에서 158개를 모두 새로 호출했다는 뜻이 아니다. 첫 실행 154회와 후속 v2 신규 시도 80회를 합하면 두 비교 단계에서 실제 신규 생성 시도는 234회다. 별도의 입력 분석·live smoke는 이 수에 포함하지 않는다. 전체 프로젝트의 모든 모델 호출 총량으로 해석하지 않는다.

기준선 재사용은 요청 내용이 완전히 같은 경우만 허용했다. 프롬프트·분석·설정이 바뀌면 같은 모델명이라는 이유로 옛 출력을 그대로 기준선으로 쓰면 안 된다. 생성 요청과 코드·KG 해시를 저장했고, 서로 다른 분석이 필요한 경우 새 조회 결과를 만들었다.

seed를 고정했어도 같은 요청에서 다른 출력이 나온 이력이 있다. 현재 최종 파일에서 동일 요청 간 차이가 없더라도 반복 실행 전체의 결정성이 입증된 것은 아니다. 서버의 실제 체크포인트·서빙 상태도 별도 기록이 필요하다. [E39]

이번 자료에는 의미·높임·문법·자연스러움의 독립 채점 완료 수치가 없다. 빈 blind_review.tsv와 개발자의 sentence_review.tsv를 구분한다. 출력이 달라진 횟수, 방언 어휘가 등장한 횟수, 참조가 채워진 횟수는 각각 진단값이다. 이를 품질 정확도로 바꾸어 발표하지 않는다.''')
h('12.1 비용을 설명할 때',2)
body('''모델 호출 비용은 지식 구축 때의 Astra 작업과 번역 시점의 Qwen 분석·생성, 임베딩 인덱스 구축을 분리해야 한다. 분석을 재사용한 실험의 생성 시간만 보고 실제 새 입력의 전체 지연이라고 말하면 안 된다. 누적 요청 시간도 병렬 실행의 벽시계 시간과 다르다. 저장된 API usage가 있으면 토큰과 시간을 집계할 수 있지만, 요금표나 모델 사용료를 확인하지 않은 상태에서 금액을 추정하지 않는다.

재현 기록에는 요청, 응답, 모델명, 설정, 시드, 실패·재시도, 입력/자료/코드 해시를 남긴다. 과거 실행의 일부 파일은 서버 원응답을 포함한 로컬 실행 폴더에 있고, 공개 보고서는 간결한 출력·프로토콜로 분리되어 있다. 인증 토큰·개인 서버 주소를 브리핑에 노출할 필요는 없다.''')

h('13. 결과와 오류 원인')
metrics=read('kg_v2/reports/translation_20261012/metrics.json')
table(['조건','성공 출력','입력과 완전히 동일'],[(k,v['successful'],v['exactly_unchanged']) for k,v in metrics['metrics']['all']['conditions'].items()])
body('''전반적으로 ‘방언 표현을 실제로 반영하는 정도’는 나아진 사례가 있다. 아직 점심을 안 먹었어→안 묵었어, 점심 먹자→점심 묵자, 어디노?→어디고?, 안 와노?→안 오노?가 그 예다. 하지만 표준어 유지가 모두 오류인 것은 아니다. 방언과 표준어가 공유하는 표현을 그대로 두는 것이 옳을 수 있다.

명백한 회귀도 있다. 집에 가는 길에 네가 좋아하는 빵을 샀어가 v2에서 집에 가노 길에 니가 좋아하노 빵을 샀어로 바뀌었다. 의문형처럼 보이는 형태를 관형절에 적용해 구조를 깨뜨린 것이다. 출력만으로 특정 KG 진술 하나가 직접 원인이라고 확정할 수는 없지만, 입력 구간을 참고 자료에 적어 두는 것만으로 적용 위치가 강제되지는 않았음을 보여준다.

높임 손실은 못 나가겠어요→몬 나가겠어, 나왔어예→나왔어 등에서 나타난다. 쉬이소→시모이소처럼 근거가 확인되지 않은 활용형도 생성됐다. 오늘→올, 먼저→앞에, 마무리→시마이 같은 대응은 문헌에 있더라도 이 문맥·말투에 적절한지 따로 판단해야 한다. 희귀한 방언형이라는 이유만으로 무조건 오답으로 단정하는 것도 피한다.''')
h('13.1 -예가 사라진 구체적인 전달 결함',2)
body('''기준 실험의 v2 패킷을 확인하면 목표 표현에 ‘예’가 들어간 ending/construction 진술이 38개 입력에서 한 번도 선택되지 않았다. KG에는 해당 진술들이 존재한다. 즉 ‘원문에 지식이 없다’는 설명만으로는 이 현상을 설명할 수 없다. [E16, E32]

조금 일찍 나왔어요의 경우 기존 KG는 -아/어예를 전달했지만, v2는 -리더·-니더·-니이더·-ㅂ니더·-ㅁ니더·-ㅁ더를 전달했다. -예와 -네예도 조건 검사를 통과할 수 있는데, 동일 구간 6개 제한과 ID 순서로 인해 제외됐다. 해예·하이소·하소를 모두 넓은 높임 집합으로 허용한 뒤 구체적인 말투 적합성을 정렬에 반영하지 않은 결함이다.

이것이 -예 계열 근거가 모델에 도달하지 않은 직접 원인이라는 점은 코드·패킷으로 확인된다. 반면 이 한 결함이 모든 높임 손실의 유일한 원인이라는 것은 아직 실험하지 않았다. 적절한 근거를 다시 전달한 뒤 같은 입력과 새 대조 문장에서 확인해야 한다.

LONG008에서는 v2 두 패킷이 17개 참조·0개 활용 관계로 완전히 같다. 힘드셨죠?는 최종 분석에서 청자 높임 unknown이라 관련 높임 어미 조회가 막힌다. 쉬세요에는 기본형 시다와 여러 높임 어미가 따로 제공됐지만 시모이소라는 완성형은 제공되지 않았다. 분석·후보 선택·활용 조합의 문제가 함께 나타난다.''')
h('13.2 관계 효과를 얼마나 확인했는가',2)
body('''활용 관계가 실제 추가된 입력은 DEV006, DEV025, LONG001, LONG002, LONG006의 5개다. 나머지 33개에서 v2 직접·관계 요청은 같았고, 최종 출력은 35/38개가 같았다. 따라서 이번 결과로 관계 구조 자체의 우월성을 주장할 수 없다.

관계가 활성화되지 않은 이유를 나누어 볼 필요가 있다. 입력이 관련 어휘를 포함하지 않거나, 대응은 있어도 해당 활용표 연결이 없거나, 과거·높임 등 복합형이라 현재 지원 범위에서 보류되거나, 예산·개수 제한으로 끝점이 빠질 수 있다. 모든 경우를 ‘KG가 쓸모없다’로 묶는 것도, 그래프가 크니 충분하다고 보는 것도 타당하지 않다.''')

h('14. 검증·평가의 경계')
table(['검증 층','확인한 것','아직 보장하지 않는 것'],[
('원본 보존','기존 보호 파일 169개 해시 유지','기존 내용 자체의 언어학적 정답'),
('문헌 범위','474쪽 페이지 대장·누락 검사','모든 정보를 완전 추출'),
('근거 문자열','인용이 실제 페이지 텍스트에 존재','인용이 진술 전체를 충분히 뒷받침'),
('의미 검토','Astra 작성·분담·교차 검토와 조정자 표본 확인','독립 전문가 전수 검증'),
('빌드/조회 테스트','초기 9개 + 문장 연결 6개 = 기준 시점 15개 통과','모든 실제 발화의 품질'),
('Neo4j','25,750노드·35,308관계의 속성까지 왕복 일치','관계가 번역에 유익하다는 인과 효과'),
('생성 실험','38개 입력 네 조건의 실제 출력 보존','미지 문장·모든 경상 방언 화자의 자연스러움'),
('개발자 출력 검토','의미·높임·문법 오류의 구체 사례','독립·블라인드 평가 점수')])
body('''선택한 지식이 맞는지와 최종 문장이 맞는지를 따로 평가해야 한다. 필요한 지식이 검색됐는데 생성이 틀린 경우와, 지식이 KG에 있는데 검색되지 않은 경우는 개선 위치가 다르다. 문헌 자체에 필요한 정보가 없는 경우도 별도로 기록해야 한다.

후속 본 실험에서는 의미 보존, 청자 높임, 주체 높임, 문법성, 자연스러움, 목표 방언 적합성을 분리해 평가할 필요가 있다. 변경량을 별도 진단으로 남길 수 있지만 목표 점수로 삼지 않는다. copy 기준선과 복수의 허용 표현을 고려하고, 지역·세대에 민감한 판단은 방언 화자 또는 관련 전문가의 평가로 확인해야 한다.

같은 지식량·같은 시작 후보에서 관계만의 효과를 보고 싶은 실험과, 실제로 더 많은 활용 지식을 제공하는 전체 시스템의 효과를 보고 싶은 실험은 별도로 설계한다. 현재 legacy/v2는 KG와 조회기가 함께 바뀌었고 direct/graph도 활용 진술의 양이 바뀐다. 현재 실험으로 각 요인의 인과 효과를 분리할 수 없다.''')

h('15. 재현과 파일 지도')
table(['찾고 싶은 내용','위치'],[
('최상위 상태','README.md / PROJECT_MAP.md / 연구미팅.md / 이관.md'),
('현재 진술 수정 원본','kg_v2/authored/*.json'),
('작성 계약·스키마','kg_v2/AUTHORING.md / fact.schema.json'),
('페이지 검토·공백','kg_v2/reviews/* / reports/page_coverage.tsv / extraction_gaps.json'),
('통합 KG','kg_v2/facts.jsonl / graph.json'),
('문장 분석·조회·생성','kg_v2/run_translation.py / sentence_retrieval.py / query.py'),
('분석 프롬프트','current_kg/experiments/run_auto_analysis_development.py'),
('현재 비교 결과','kg_v2/reports/translation_20261012/'),
('수정 전 비교 보존','kg_v2/reports/translation_20261012_first/'),
('원요청/응답·실행 스냅샷','kg_v2/local/translation_* 및 current_kg/experiments/archive/runs/'),
('Neo4j 그림과 검사','kg_v2/reports/neo4j/'),
('기존 KG','current_kg/grammar/'),
('이전 연구 이력','archive/v1_retrieval/ / archive/v2_failed/ / current_kg/docs/history/')])
body('''facts.jsonl과 graph.json을 수동 편집하기보다 authored의 원본 진술을 고치고 다시 빌드한다. 문장 조회기만 바꾸는 경우 KG 원본 자체가 달라질 필요는 없다. 새 실행은 새 폴더에 저장하고 기존 결과를 덮어쓰지 않는다. Git 기록은 3c957a3에 전체 원문 KG v2·병렬 검토, 8eee9ce에 문장 생성 연결과 Neo4j 그림이 반영되어 있다.''')
code('python kg_v2/extract_sources.py\n.venv/retrieval/bin/python kg_v2/build.py --require-full\n.venv/retrieval/bin/python -m unittest discover -s kg_v2 -p \'test_*.py\' -v\npython kg_v2/integration_audit.py\npython kg_v2/query.py --lemma 먹다 --ending=-고\n\n# 새 입력은 기존에 없는 출력 폴더를 지정\n.venv/retrieval/bin/python kg_v2/run_translation.py translate \\\n  --text \'나는 배가 고파서 먼저 밥을 먹고 있을게.\' \\\n  --folder kg_v2/local/new_live_run\n\n# 비교 실행: 같은 새 폴더로 준비→생성→내보내기\n.venv/retrieval/bin/python kg_v2/run_translation.py prepare --folder kg_v2/local/new_comparison\n.venv/retrieval/bin/python kg_v2/run_translation.py generate --folder kg_v2/local/new_comparison\n.venv/retrieval/bin/python kg_v2/run_translation.py export --folder kg_v2/local/new_comparison')
body('''재현에는 로컬 원본 PDF, Python 의존성, Qwen 서버 접근 설정이 필요하다. Neo4j 시각화는 별도로 서버와 적재 절차가 필요하다. 인증정보는 기존 로컬 설정에서 읽고 출력하지 않는다. 현재 보고서 폴더의 프로토콜·해시를 함께 보아야 같은 조건인지 판단할 수 있다.''')

h('16. 선행연구와 논문 구성')
body('''방법론 문헌은 저장소에 11편의 검토 기록이 있다. 이것은 보유 자료에 대한 연구 배경 검토이며 체계적 문헌조사나 모든 논문의 재현 실험은 아니다. 아래 설명은 그 기존 검토 기록이 우리 설계에 어떤 질문을 남겼는지를 정리한 것이다. 논문에 수치를 직접 인용할 때는 출판본과 해당 표·조건을 다시 확인한다. [E10]''')
table(['관련 연구 묶음','우리 연구에 필요한 질문'],[
('EDC / 온톨로지 학습 기반 KG 구축','추출한 표현을 어떻게 통합하고 출처·의미 차이를 보존할 것인가.'),
('One Grammar Book / LingoLLM','문법 설명·예문·사전·형태 분석 중 실제 도움이 되는 요소는 무엇인가.'),
('Can LLMs Help Create Grammar?','LLM 작성 초안의 형식적 적합성과 언어학적 타당성을 어떻게 구분할 것인가.'),
('GRAMMAR-LLM','출력 형식을 지키게 하는 것과 자연스러운 방언을 생성하는 것은 어떻게 다른가.'),
('KAPING / SURGE / GraphRAG','관련 지식 검색과 생성에서의 활용을 어떻게 분리 평가할 것인가. 과제·학습 조건이 다름.'),
('DIA-REFINE','표준어 복사와 분류기 피드백의 효과를 어떻게 독립적인 품질 평가와 구분할 것인가.'),
('방언 정규화·피벗 NMT','방언→표준어 정규화 성과를 표준어→방언 생성 성과로 잘못 인용하지 않는가.')])
body('''논문 방법 절에는 자료·작성 계약·진술 스키마·조건 주체·그래프 컴파일·입력 분석·조회·생성 프로토콜을 기술할 수 있다. 실험 절에는 개발 실험임을 명시한 결과와 실패 분석을 제시할 수 있다. 그러나 독립 평가가 없는 상태에서 KG의 일반적인 우수성을 결론으로 먼저 적어서는 안 된다.

가능한 현재 기여는 출처와 적용 범위를 추적하는 방언 지식 표현, 표준어 기본형→방언 어휘→문헌 활용형의 연결, 모델 호출·조회·출력의 재현 가능한 기록, 그리고 실제 실패 위치의 분해다. 이 기여도 단순히 KG를 만들었다는 사실만으로 새로움이 확정되는 것은 아니며, 관련 연구와의 구조·평가 차이를 논문에서 명확히 해야 한다.''')

h('17. 다음 단계와 브리핑 질의응답')
body('''우선순위 1은 현재 확인된 전달 결함을 고치는 것이다. 높임을 단순 polite/non_polite로만 다루지 않고 원문에 드러난 말투를 보수적으로 추정하여 해예·하이소·하소 계열 후보의 순서를 정한다. 적합한 후보가 ID 순서 때문에 밀리지 않도록 하고, 후보 부족이나 분석 미확인을 기록한다. 이는 특정 문장에 정답을 하드코딩하거나 모든 문장에 예를 붙이라는 뜻이 아니다.

우선순위 2는 같은 38개 입력의 조회 전후를 비교해 필요한 높임 근거가 도달하는지 확인하는 것이다. 질문/평서/요청과 반말/해요형/격식형 등 새로운 대조 문장도 확인한다. 생성까지 재실행했다면 그 결과는 새 실험으로 보존하고, 조회만 검사했으면 번역 품질 개선이라고 보고하지 않는다.

우선순위 3은 적용 위치·높임·시제·의미 검증을 별도 설계하고, 고정된 개발 종료 기준 이후 독립 평가를 하는 것이다. 입력 분석·후보 검색·생성·후처리 효과를 한 번에 모두 바꿔 인과 해석을 어렵게 하지 않는다. 기존 38개는 이미 오류 수정에 사용했으므로 독립 시험셋으로 되돌릴 수 없다.''')
qas=[
('KG는 Astra, 번역은 Qwen이라고 말하면 맞는가?','핵심 역할 요약으로는 맞다. 다만 Qwen 입력 분석·문헌 대응·관련성 검토와 E5 임베딩, 역사적 Claude 실험을 추가로 설명해야 한다.'),
('Astra 세 개가 각각 KG 전체를 만들었나?','아니다. 전체 문헌 초안 뒤 남은 162쪽의 항목을 세 묶음으로 분담했다. 조정자가 공통 스키마·통합·근거 검증을 맡았다.'),
('왜 v2가 두 줄인가?','같은 v2 지식에서 완성 활용 지식과 관계를 추가 전달하는 효과를 보기 위한 직접/관계 조건이다. LONG008은 추가 경로가 없어 두 입력 패킷까지 같다.'),
('현재 모델이 Neo4j에 질문하는가?','문장 런타임은 Python이 로컬 그래프를 조회한다. Neo4j는 별도 적재·시각화·왕복 검증 경로다.'),
('모든 30문장이 같은 지식을 받았나?','아니다. 각 문장 안에서 직접 조건과 관계 조건의 선택이 같았다는 의미였다. 문장 간 조회 결과가 모두 같다는 뜻이 아니다.'),
('왜 문헌 전체를 읽고도 번역이 나쁜가?','지식의 존재, 검색 도달, 조건 적합성, 생성 조합, 결과 검증은 다른 단계다. -예 누락은 지식 부족이 아니라 선택 단계의 결함을 보여준다.'),
('그래프가 더 커졌으니 좋아진 것 아닌가?','노드 수는 표현 방식과 분해 단위에 좌우된다. 적절한 후보의 도달과 최종 품질을 따로 확인해야 한다.'),
('일부 문장이 표준어 그대로면 실패인가?','공유 표현은 유지가 정답일 수 있다. 다만 체계적으로 필요한 방언형을 못 쓰는 현상은 구체적인 문장과 지식 도달을 함께 보아야 한다.'),
('현재 가장 분명한 성과는?','새 KG의 출처·조건·활용 경로를 구축하고 실제 분석·생성에 연결했으며, 성공과 실패를 요청·출력 수준에서 추적할 수 있다.'),
('아직 가장 부족한 것은?','원문 말투에 맞는 선택과 활용 통제, 의미·높임 보존 검증, 독립 문장과 방언 화자 평가다.'),
('KG v2가 기존보다 낫다는 결론인가?','어휘·일부 의문형 개선과 문법·높임 회귀가 함께 관찰됐다. 전체 품질 우위나 관계의 독립 효과를 결론내릴 수 없다.'),
('현재 결과로 발표해도 되는가?','개발 연구의 과정·구현·실패 진단으로 설명할 수 있다. 완료한 단계와 품질 검증이 남은 단계를 명확히 구분하면 된다.')]
for q,a in qas:
    h(q,2);p(a)

h('부록 A. 모델 프롬프트와 입출력 계약')
h('A.1 Astra 문헌 작성·검토: 고정 HTTP 프롬프트가 아닌 작업 계약',2)
body('''입력은 담당 문헌·페이지, 작성 계약, JSON 스키마, 기존 진술·gap 대장과 수정 요청이다. 에이전트가 원문을 읽고 필요하면 표 이미지를 확인한 뒤 authored와 reviews를 작성한다. 출력은 번역문이 아니라 진술·검토 JSON과 완료 보고다.

다음은 AUTHORING.md와 검토 보고서를 바탕으로 재구성한 지시 구조이며, 당시 spawn_agent.message의 원문이 아니다: “담당 문헌의 남은 항목을 원문과 대조하고 공통 스키마를 유지하라. 조건의 판정 주체와 문헌 범위를 보존하라. 미해석 표를 복원하지 말고 한계로 기록하라. 담당 authored/reviews와 완료 보고를 작성하고 통합 산출물은 조정자가 빌드하도록 하라.”

모델 ID와 실제 분담은 기록으로 확인했지만, 문헌 작성 전체를 완전히 같은 단일 프롬프트와 토큰 수로 재현할 수 있는 자동 추출 API 파이프라인은 아니다. 이 한계를 방법 절에 명시해야 한다.''')
h('A.2 Qwen 입력 분석: 실제 시스템 지시문',2)
sys.path.insert(0,str(ROOT/'current_kg/experiments'))
from run_auto_analysis_development import analysis_system, raw_schema
analysis_prompt=analysis_system()
instructions,extra=analysis_prompt.split('입력 특징 정의:',1)
code(instructions.strip())
body('''위 지시문 뒤에는 analysis_contract.json의 입력 특징 정의와 raw_schema()의 JSON 스키마가 실제로 이어 붙는다. user 메시지는 표준어 문자열이다. 요청은 temperature=0, top_p=1, seed=20260923, max_tokens=4200, enable_thinking=false, response_format=json_object다. 최대 두 번 시도하며 잘못된 JSON·원문 결속 실패를 기록한다. 아래가 실제 출력 스키마다. 방언 정답은 포함하지 않는다.''')
code(raw_schema())
(OUT/'analysis_system_prompt.txt').write_text(analysis_prompt+'\n')
h('A.3 Qwen 문헌 1차 대응: 실제 시스템 지시문',2)
code(literal_assignment('current_kg/experiments/review_source_inventory.py','review_batch','system'))
body('''user 메시지는 [{id, source_id, pdf_page, text, kg_candidates:[{id, description}]}] 형태다. temperature=0, top_p=1, seed=20260923, max_tokens=6500, JSON 응답을 사용한다. 출력은 units 배열이고 kind/summary/examples/proposed_record_ids/needs_visual_review/limits를 갖는다. 인용문은 원문에 실제 존재해야 하며 제공하지 않은 KG ID는 격리한다. 이 역할은 1차 대응 제안이지 지식의 최종 승인이나 번역이 아니다.''')
h('A.4 Qwen 검색 관련성 검토: 실제 시스템 지시문',2)
code(literal_assignment('current_kg/experiments/source_retrieval.py','semantic_review','system'))
body('''user 메시지는 {input, input_units, candidates}다. 문자 검색으로 상위 후보를 만든 후 모델에 관련 항목을 최대 6개 고르게 한다. 출력은 selected 배열의 id/input_unit_id/evidence_quote/reason이다. temperature=0, top_p=1, seed=20260923, max_tokens=2400, JSON 응답이다. 관련성이 있다고 판단된 것과 적용 조건이 충족된 것은 다르다. 설명 reason은 감사에 남기되 근거인 것처럼 생성 자료에 추가하지 않는다.''')
h('A.5 Qwen 방언 생성: 실제 공통 프롬프트',2)
code(protocol['config']['system_prompt'])
code(literal_assignment('kg_v2/run_translation.py','make_request','instruction'))
body('''위 user 지시문 뒤에 {standard, analysis, source_morphology, references, graph_paths} JSON이 이어진다. 네 조건의 시스템 지시·user 공통 지시·모델·시드·생성 설정은 같다. 조건별로 references와 graph_paths를 바꾼다. 출력은 JSON이 아니라 변환한 발화만 있는 문자열이다. 서버 응답에서 문자열을 꺼내 원형 그대로 보존한다.''')
h('A.6 임베딩·과거 모델 사용의 입력과 출력',2)
body('''E5 입력은 passage: 문서조각과 query: 질의문이다. 평균 풀링 후 L2 정규화한 벡터를 만들고 내적으로 순위를 계산한다. 이는 자연어 응답을 생성하는 채팅 프롬프트가 아니다.

초기 S1은 시나리오 키워드와 지역·지식/예시 블록을 넣어 상담원/민원인 대화를 직접 생성했다. S2는 표준어 대화를 생성한 뒤 민원인 발화를 검색 질의로 사용하여 방언으로 변환했다. 과거 Claude 교란에서는 생성된 대화를 표준어화하거나 다른 권역 형태로 바꿔 지표가 반응하는지 보았다. 이 과정의 출력은 각각 대화 텍스트 또는 교란/평가 JSON이다. 현재의 문장 분석 JSON과 혼동하지 않는다. 과거의 모든 모델 설정을 현재 Qwen 설정으로 통일해 적지 않는다.''')
h('A.7 초기 S1/S2·키워드 추출 프롬프트 원문',2)
body('''아래는 archive에 보존된 템플릿이다. 중괄호는 시나리오·권역·검색 지식 등으로 채워지는 자리다. 현재 38개 실험에서 사용한 프롬프트가 아니다. 실행 코드의 조건별 추가 블록과 실제 요청을 함께 보아야 과거의 완전한 입력을 재구성할 수 있다. 키워드 추출 템플릿이 있다는 사실만으로 모든 과거 생성물에 해당 평가가 실행됐다고 가정하지 않는다.''')
for name,label in [('prompt_template.txt','S1 직접 방언 대화 생성'),('prompt_standard.txt','S2의 표준어 대화 생성'),('convert_template.txt','S2의 방언 변환'),('extract_prompt.txt','초기 5키워드 추출 평가')]:
    h(label,3)
    p('원본: archive/v1_retrieval/prompts/'+name,'Caption')
    code((ROOT/'archive/v1_retrieval/prompts'/name).read_text())
h('A.8 과거 Claude 교란·쌍대 평가 템플릿',2)
body('''아래는 실제 보존 코드의 템플릿 함수다. {dialogue}, {src}, {tgt}, {region}, {t1}, {t2}는 각 실행의 입력으로 채운다. 교란 결과 파일은 존재한다. 쌍대 평가 코드는 존재하지만 이 문서는 모든 과거 실행·기본 모델명을 실제 호출 기록으로 확인한 것은 아니다. 코드 기본값 claude-opus-4-8을 현재 독립 평가에 사용한 모델이라고 쓰지 않는다. 38문장의 독립 채점은 완료되지 않았다.''')
for path,fn in [('archive/v1_retrieval/src/eval/llm_perturb_sensitivity.py','prompt_standardize'),('archive/v1_retrieval/src/eval/llm_perturb_sensitivity.py','prompt_crossregion'),('archive/v1_retrieval/src/eval/judge_conditions.py','prompt_pair')]:
    code(source_excerpt(path,fn))

h('부록 B. KG·분석·조회 실물 예시')
facts={f['id']:f for f in (json.loads(l) for l in (ROOT/'kg_v2/facts.jsonl').read_text().splitlines())}
h('B.1 실제 어휘 대응 진술',2)
code(facts['NAMU_GRAMMAR_20260908:v2:full:1057'])
h('B.2 실제 활용 진술',2)
code(facts['NAMU_GRAMMAR_20260908:v2:full:0239'])
h('B.3 실제 입력 분석 예: DEV026',2)
analysis=read('kg_v2/reports/translation_20261012/analysis.json')
example=next(i for i in analysis['items'] if i['id']=='DEV026')
code({k:example[k] for k in ['id','input','preservation_note','clauses']})
h('B.4 실제 생성 참조 항목 예',2)
retrieval=read('kg_v2/reports/translation_20261012/retrieval.json')
r=next(r for r in retrieval if r['id']=='LONG001')
lex=next(f for f in r['packets']['v2_graph']['references'] if f['id']=='NAMU_GRAMMAR_20260908:v2:full:1057')
code(lex)
body('''input_binding이 원문 어느 부분에 이 지식을 연결했는지 보여준다. checks는 조건 판정이고 scope/constraints는 범위다. status가 conditional_reference_not_approved_for_generation이라는 점에 주목한다. 이 항목을 보냈다고 모델이 해당 부분에만 적용하도록 강제되는 것은 아니다.''')
h('B.5 그래프에 존재하는 관계 유형과 수',2)
table(['출발','관계','도착','수'],[(e['from'],e['relation'],e['to'],e['count']) for e in counts['edges']])
h('B.6 진술 유형·근거 상태 분포',2)
table(['진술 유형','수'],list(summary['fact_types'].items()))
table(['근거/주장 상태','수'],list(summary['claim_statuses'].items()))

h('부록 C. 38개 입력의 네 조건 출력 전체')
body('''다음은 translation_20261012/outputs.json의 152개 성공 출력을 그대로 수록한 것이다. 띄어쓰기·어휘·오류를 문서 작성자가 교정하지 않았다. legacy_graph는 기존 current_kg의 조회 결과를 공통 생성 프롬프트에 넣은 조건이다. 10월 8일 파일의 문자열을 단순 복사한 조건이라는 뜻은 아니다. 아래 출력은 문서 작성 이후의 높임 우선순위 수정 결과가 아니다.''')
outputs=read('kg_v2/reports/translation_20261012/outputs.json')
grouped=defaultdict(dict)
for row in outputs:grouped[row['id']][row['condition']]=row
for sid,rows in grouped.items():
    h(sid+' · '+rows['no_kg']['input'],2)
    for mode in ['no_kg','legacy_graph','v2_direct','v2_graph']:
        par=p('');run=par.add_run(mode+': ');run.bold=True;par.add_run(rows[mode]['output'])

h('부록 D. 근거 파일과 무결성')
body('''본문의 근거 번호는 아래 파일을 가리킨다. 상대 경로는 저장소 루트 기준이다. evidence_inventory.json에는 문서 정리 시작 시점의 전체 SHA-256을 기록했다. 문서 이후 코드 수정은 별도 부록과 변경 보고서로 구분한다. 비공개 인증 설정이나 원문 PDF 전체를 이 문서에 복사하지 않았다.''')
manifest=read('reports/briefing_20261012/evidence_inventory.json')
for e in manifest:
    p(e['id']+'  '+e['path'])
    p('SHA-256 '+e['sha256'],'Caption')
h('D.1 이번 정리의 확인 범위',2)
body('''현재 코드의 모델 호출·입출력·조회·생성 구조, 최신 저장 출력·프로토콜·검토·그래프 통계를 대조했다. 역사적 연구 결과는 해당 시점의 감사·서술 기록을 근거로 요약했으며 과거 실험 전체를 재실행하지 않았다. 세 Astra 작업자의 역할·모델명은 세션 메타데이터와 보고서를 대조했다. 정확한 작업 메시지 원문은 확인 불가한 범위를 명시했다.

보고서 수치 중 build_summary.json의 translation_quality=not evaluated나 초기 verification의 translation_evaluation=not run은 지식 구축 당시의 상태다. 이후 문장 실험을 실행한 translation_20261012 보고서를 우선한다. 오래된 단계의 상태 문자열을 최신 전체 상태로 오해하지 않는다.''')

h('부록 E. 문서 작성 이후 높임 조회 수정')
change=OUT/'politeness_followup.json'
if change.exists():
    record=json.loads(change.read_text())
    for title,text in record['sections']:
        h(title,2);body(text)
    if 'metrics' in record:table(['검사 항목','결과'],record['metrics'])
else:
    body('''본문과 38문장 부록을 먼저 완성한 뒤, 원문 말투에 맞는 높임 후보 우선순위 로직을 수정할 예정이다. 현재 이 초안에는 수정 전의 실험과 진단만 포함한다. 수정 및 검증이 끝나면 이 부록에 구현 내용·조회 전후·테스트·번역 재실행 여부를 추가한다.''')

dest=OUT/'경상도_방언_KG_연구_상세브리핑_20261012.docx'
doc.save(dest)
check=Document(dest)
alltext='\n'.join(p.text for p in check.paragraphs)
assert len(outputs)==152 and len(grouped)==38
assert all(row['output'] in alltext for row in outputs)
validation={'file':dest.name,'bytes':dest.stat().st_size,'paragraphs':len(check.paragraphs),
            'tables':len(check.tables),'embedded_images':len(check.inline_shapes),
            'source_outputs_included':152,'input_groups':38,'headings':sum(p.style.name.startswith('Heading') for p in check.paragraphs),
            'word_processor_rendering':'not yet rendered; OOXML reopened and content checked',
            'sha256':hashlib.sha256(dest.read_bytes()).hexdigest()}
(OUT/'document_validation.json').write_text(json.dumps(validation,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(validation,ensure_ascii=False,indent=2))
